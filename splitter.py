from __future__ import annotations

from collections import OrderedDict
from copy import copy
from dataclasses import dataclass, field
from decimal import Decimal, ROUND_HALF_UP
import os
from pathlib import Path
import re
import shutil
from typing import Callable, Iterable

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties


ProgressCallback = Callable[[int, str], None]
OrderProgressCallback = Callable[[int, int], None]


class SplitterError(Exception):
    """An input workbook cannot be processed safely."""


TWO_DECIMALS = Decimal("0.01")


def _round_two(value: float) -> float:
    """Round like Excel ROUND(value, 2), not Python's bankers rounding."""
    return float(
        Decimal(str(value)).quantize(TWO_DECIMALS, rounding=ROUND_HALF_UP)
    )


def _allocate_two(total: float, weights: Iterable[float]) -> list[float]:
    """Allocate a two-decimal total and put the cent remainder on the last row."""
    weight_values = [Decimal(str(value)) for value in weights]
    if not weight_values:
        return []
    weight_total = sum(weight_values)
    if weight_total == 0:
        raise ValueError("分摊权重合计不能为 0")
    target = Decimal(str(_round_two(total)))
    allocated = Decimal("0")
    result: list[float] = []
    for index, weight in enumerate(weight_values):
        amount = (
            target - allocated
            if index == len(weight_values) - 1
            else (target * weight / weight_total).quantize(
                TWO_DECIMALS,
                rounding=ROUND_HALF_UP,
            )
        )
        allocated += amount
        result.append(float(amount))
    return result


@dataclass
class SplitStats:
    orders: int = 0
    rows: int = 0
    matched_rows: int = 0
    unmatched_rows: int = 0
    exceptional_orders: int = 0
    warnings: list[str] = field(default_factory=list)
    output_path: str = ""


@dataclass
class RatioUpdateStats:
    mode: str
    added_rows: int
    source_rows: int
    unique_items: int


@dataclass(frozen=True)
class RatioEntry:
    price: float
    ratio: float
    row: int


@dataclass
class RatioData:
    prices: OrderedDict[str, float]
    pairs: dict[tuple[str, str], RatioEntry]
    source_rows: int
    has_parent_ratios: bool


@dataclass(frozen=True)
class SalesColumns:
    order_number: int
    logistics_number: int
    web_order_number: int
    receivable: int
    item_code: int
    quantity: int
    unit_price: int
    original_amount: int
    note: int
    marker: int | None = None


@dataclass(frozen=True)
class ResultColumns:
    bundle_price: int
    ratio: int
    split_unit_price: int
    split_amount: int
    allocated_amount: int
    source_unit_price: int | None = None
    fee: int | None = None
    final_amount: int | None = None

    def all(self) -> tuple[int, ...]:
        return tuple(
            column
            for column in (
                self.bundle_price,
                self.ratio,
                self.split_unit_price,
                self.split_amount,
                self.allocated_amount,
                self.source_unit_price,
                self.fee,
                self.final_amount,
            )
            if column is not None
        )


@dataclass(frozen=True)
class ResultLayout:
    sales: SalesColumns
    result: ResultColumns
    source_columns: tuple[int, ...]


SALES_REQUIRED_HEADERS = {
    "订单编号",
    "物流单号",
    "网店订单号",
    "应收合计",
    "货品编号",
    "数量",
    "单价",
    "金额",
    "备注",
}
# Kept as a public compatibility alias for existing integrations and tests.
SALES_HEADERS = SALES_REQUIRED_HEADERS

RATIO_SCHEMAS = (
    (
        {"单品编号", "金额", "分摊比例", "单价"},
        "单品编号",
        "单价",
        "",
        "",
        "Sheet2",
    ),
    (
        {"母件编号", "编号", "执行价格", "分摊金额", "分摊比例"},
        "编号",
        "执行价格",
        "母件编号",
        "分摊比例",
        "sheet1",
    ),
)
SPLIT_RESULT_HEADERS = ("组合装单价", "占比", "单价", "金额")
AMOUNT_RESULT_HEADERS = ("单价2", "手续费1%", "最终金额", "摊后金额")


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _header_name(value: object) -> str:
    return re.sub(r"\s+", "", _text(value).replace("\u3000", ""))


def _code(value: object) -> str:
    return _text(value).upper()


def _number(value: object, label: str) -> float:
    if isinstance(value, bool):
        raise ValueError(f"{label}不是数字")
    if isinstance(value, (int, float)):
        return float(value)
    text = _text(value).replace(",", "")
    if not text:
        raise ValueError(f"{label}为空")
    try:
        return float(text)
    except ValueError as exc:
        raise ValueError(f"{label}不是数字：{value}") from exc


def _find_sheet(workbook, required: set[str], preferred: str | None = None):
    candidates = list(workbook.worksheets)
    if preferred in workbook.sheetnames:
        candidates.remove(workbook[preferred])
        candidates.insert(0, workbook[preferred])

    matches = []
    for sheet in candidates:
        for row_number in range(1, min(sheet.max_row, 5) + 1):
            headers = {
                _header_name(sheet.cell(row_number, col).value)
                for col in range(1, sheet.max_column + 1)
            }
            if required.issubset(headers):
                matches.append((sheet, row_number))
                break
    if matches:
        return max(matches, key=lambda match: (match[0].title == preferred, match[0].max_row))
    return None, None


def _header_map(sheet, header_row: int) -> dict[str, int]:
    result: dict[str, int] = {}
    for column in range(1, sheet.max_column + 1):
        name = _header_name(sheet.cell(header_row, column).value)
        if name and name not in result:
            result[name] = column
    return result


def _sales_columns(sheet, header_row: int) -> SalesColumns:
    headers = _header_map(sheet, header_row)
    missing = SALES_REQUIRED_HEADERS.difference(headers)
    if missing:
        raise SplitterError(f"销售表缺少字段：{'、'.join(sorted(missing))}")
    amount_columns = [
        column
        for column in range(1, sheet.max_column + 1)
        if _header_name(sheet.cell(header_row, column).value) == "金额"
    ]
    return SalesColumns(
        order_number=headers["订单编号"],
        logistics_number=headers["物流单号"],
        web_order_number=headers["网店订单号"],
        receivable=headers["应收合计"],
        item_code=headers["货品编号"],
        quantity=headers["数量"],
        unit_price=headers["单价"],
        original_amount=amount_columns[-1],
        note=headers["备注"],
        marker=headers.get("标记"),
    )


def _ratio_sheet_info(workbook):
    for (
        required,
        code_header,
        price_header,
        mother_header,
        ratio_header,
        preferred,
    ) in RATIO_SCHEMAS:
        sheet, header_row = _find_sheet(workbook, required, preferred)
        if sheet is not None:
            return (
                sheet,
                header_row,
                code_header,
                price_header,
                mother_header,
                ratio_header,
            )
    return None, None, "", "", "", ""


def _load_ratio_data(ratio_path: str | Path) -> RatioData:
    path = Path(ratio_path)
    if path.suffix.lower() != ".xlsx":
        raise SplitterError("组合装拆分占比表必须是 .xlsx 文件。")

    try:
        values_book = load_workbook(path, data_only=True, read_only=True, keep_links=False)
        formulas_book = load_workbook(path, data_only=False, read_only=True, keep_links=False)
    except Exception as exc:
        raise SplitterError(f"无法打开组合装拆分占比表：{exc}") from exc

    try:
        (
            value_sheet,
            header_row,
            code_header,
            price_header,
            mother_header,
            ratio_header,
        ) = _ratio_sheet_info(values_book)
        if value_sheet is None:
            raise SplitterError(
                "占比表中未找到可识别的工作表。支持旧版“单品编号、金额、分摊比例、单价”，"
                "或新版“母件编号、编号、执行价格、分摊金额、分摊比例”。"
            )
        formula_sheet = formulas_book[value_sheet.title]
        headers = _header_map(value_sheet, header_row)
        code_col = headers[code_header]
        price_col = headers[price_header]
        mother_col = headers[mother_header] if mother_header else None
        ratio_col = headers[ratio_header] if ratio_header else None

        prices: OrderedDict[str, float] = OrderedDict()
        pairs: dict[tuple[str, str], RatioEntry] = {}
        missing_cache: list[tuple[int, str]] = []
        source_rows = 0
        value_rows = value_sheet.iter_rows(min_row=header_row + 1, values_only=True)
        formula_rows = formula_sheet.iter_rows(min_row=header_row + 1, values_only=True)
        for row, (values, formulas) in enumerate(
            zip(value_rows, formula_rows), start=header_row + 1
        ):
            item_code = _code(values[code_col - 1])
            if not item_code:
                continue
            cached_value = values[price_col - 1]
            # 合并多个导出文件时，数据区中可能再次出现完整表头。
            if item_code == _code(code_header) and _text(cached_value) == price_header:
                continue
            source_rows += 1
            formula_value = formulas[price_col - 1]
            if cached_value is None and isinstance(formula_value, str) and formula_value.startswith("="):
                missing_cache.append((row, price_header))
                continue
            try:
                price = _number(cached_value, f"占比表第 {row} 行{price_header}")
            except ValueError as exc:
                raise SplitterError(str(exc)) from exc
            prices.setdefault(item_code, price)

            if mother_col is None or ratio_col is None:
                continue

            mother_code = _code(values[mother_col - 1])
            if not mother_code:
                raise SplitterError(f"占比表第 {row} 行母件编号为空。")
            ratio_value = values[ratio_col - 1]
            ratio_formula = formulas[ratio_col - 1]
            if ratio_value is None and isinstance(ratio_formula, str) and ratio_formula.startswith("="):
                missing_cache.append((row, ratio_header))
                continue
            try:
                ratio = _number(ratio_value, f"占比表第 {row} 行{ratio_header}")
            except ValueError as exc:
                raise SplitterError(str(exc)) from exc

            key = (mother_code, item_code)
            existing = pairs.get(key)
            if existing is not None:
                if (
                    abs(existing.price - price) > 1e-9
                    or abs(existing.ratio - ratio) > 1e-9
                ):
                    raise SplitterError(
                        f"占比表母件 {mother_code}、子件 {item_code} 在第 "
                        f"{existing.row} 行和第 {row} 行存在不同的执行价格或分摊比例。"
                    )
                continue
            pairs[key] = RatioEntry(price, ratio, row)

        if missing_cache:
            preview = "、".join(
                f"{row} 行{header}" for row, header in missing_cache[:8]
            )
            raise SplitterError(
                f"占比表公式没有可读取的计算结果（第 {preview}）。"
                "请先用 Excel 打开该文件，完成计算并保存后再试。"
            )
        if not prices:
            raise SplitterError("占比表没有可用的子件编号和拆分单价。")
        if mother_col is not None and not pairs:
            raise SplitterError("占比表没有可用的母件编号、子件编号和分摊比例。")
        return RatioData(
            prices=prices,
            pairs=pairs,
            source_rows=source_rows,
            has_parent_ratios=mother_col is not None,
        )
    finally:
        values_book.close()
        formulas_book.close()


def ratio_file_info(ratio_path: str | Path) -> tuple[int, int]:
    ratio_data = _load_ratio_data(ratio_path)
    return len(ratio_data.prices), ratio_data.source_rows


def update_ratio_data(
    current_path: str | Path,
    incoming_path: str | Path,
    mode: str,
) -> RatioUpdateStats:
    current = Path(current_path)
    incoming = Path(incoming_path)
    if mode not in {"append", "replace"}:
        raise SplitterError("基础数据更新方式必须是新增或覆盖。")

    incoming_items, incoming_rows = ratio_file_info(incoming)
    current.parent.mkdir(parents=True, exist_ok=True)
    temporary = current.with_name(current.stem + ".updating.xlsx")

    if mode == "replace":
        shutil.copyfile(incoming, temporary)
        os.replace(temporary, current)
        return RatioUpdateStats(mode, incoming_rows, incoming_rows, incoming_items)

    ratio_file_info(current)
    try:
        base_book = load_workbook(current, data_only=False, keep_links=False)
        incoming_book = load_workbook(incoming, data_only=True, keep_links=False)
    except Exception as exc:
        raise SplitterError(f"无法打开基础数据文件：{exc}") from exc

    try:
        (
            base_sheet,
            base_header_row,
            base_code_header,
            base_price_header,
            _,
            _,
        ) = _ratio_sheet_info(base_book)
        (
            new_sheet,
            new_header_row,
            new_code_header,
            new_price_header,
            _,
            _,
        ) = _ratio_sheet_info(incoming_book)
        if base_sheet is None or new_sheet is None:
            raise SplitterError("新增文件中未找到可识别的组合装数据工作表。")

        base_headers = _header_map(base_sheet, base_header_row)
        new_headers = _header_map(new_sheet, new_header_row)
        base_mother_header = "母件编号" if "母件编号" in base_headers else "组合装编号"
        new_mother_header = "母件编号" if "母件编号" in new_headers else "组合装编号"
        base_mother_col = base_headers.get(base_mother_header)
        new_mother_col = new_headers.get(new_mother_header)
        base_code_col = base_headers[base_code_header]
        new_code_col = new_headers[new_code_header]
        new_price_col = new_headers[new_price_header]

        existing: set[tuple[str, str]] = set()
        for values in base_sheet.iter_rows(min_row=base_header_row + 1, values_only=True):
            code = _code(values[base_code_col - 1])
            if not code or code == _code(base_code_header):
                continue
            mother = _code(values[base_mother_col - 1]) if base_mother_col else ""
            existing.add((mother, code))

        style_row = min(base_header_row + 1, base_sheet.max_row)
        added_rows = 0
        for values in new_sheet.iter_rows(min_row=new_header_row + 1, values_only=True):
            code = _code(values[new_code_col - 1])
            if not code or code == _code(new_code_header):
                continue
            mother = _code(values[new_mother_col - 1]) if new_mother_col else ""
            key = (mother, code)
            if key in existing:
                continue

            target_row = base_sheet.max_row + 1
            for header, target_col in base_headers.items():
                source_header = header
                if header == base_code_header:
                    source_header = new_code_header
                elif header == base_price_header:
                    source_header = new_price_header
                elif header == base_mother_header:
                    source_header = new_mother_header
                source_col = new_headers.get(source_header)
                value = values[source_col - 1] if source_col else None
                target = base_sheet.cell(target_row, target_col, value)
                template = base_sheet.cell(style_row, target_col)
                if template.has_style:
                    target._style = copy(template._style)
                target.number_format = template.number_format
                target.alignment = copy(template.alignment)
            existing.add(key)
            added_rows += 1

        if base_sheet.auto_filter.ref:
            start = base_sheet.auto_filter.ref.split(":", 1)[0]
            base_sheet.auto_filter.ref = f"{start}:{get_column_letter(base_sheet.max_column)}{base_sheet.max_row}"
        for table in base_sheet.tables.values():
            start = table.ref.split(":", 1)[0]
            table.ref = f"{start}:{get_column_letter(base_sheet.max_column)}{base_sheet.max_row}"

        base_book.save(temporary)
    finally:
        base_book.close()
        incoming_book.close()

    os.replace(temporary, current)
    unique_items, source_rows = ratio_file_info(current)
    return RatioUpdateStats(mode, added_rows, source_rows, unique_items)


def _copy_column_layout(
    sheet,
    old_layout: dict[int, dict[str, object]],
    mapping: dict[int, int],
    result_columns: ResultColumns,
) -> None:
    for old_col, new_col in mapping.items():
        props = old_layout.get(old_col, {})
        dim = sheet.column_dimensions[get_column_letter(new_col)]
        for name, value in props.items():
            setattr(dim, name, value)

    widths = {
        result_columns.bundle_price: 16,
        result_columns.ratio: 15,
        result_columns.split_unit_price: 16,
        result_columns.split_amount: 16,
        result_columns.allocated_amount: 16,
        result_columns.source_unit_price: 13,
        result_columns.fee: 13,
        result_columns.final_amount: 13,
    }
    for col, width in widths.items():
        if col is None:
            continue
        sheet.column_dimensions[get_column_letter(col)].width = width


def _capture_column_layout(sheet) -> dict[int, dict[str, object]]:
    layout: dict[int, dict[str, object]] = {}
    for col in range(1, sheet.max_column + 1):
        dim = sheet.column_dimensions[get_column_letter(col)]
        layout[col] = {
            "width": dim.width,
            "hidden": dim.hidden,
            "outlineLevel": dim.outlineLevel,
            "bestFit": dim.bestFit,
            "collapsed": dim.collapsed,
        }
    return layout


def _style_new_columns(
    sheet,
    header_row: int,
    sales_columns: SalesColumns,
    result_columns: ResultColumns,
) -> None:
    for row in range(1, sheet.max_row + 1):
        source = sheet.cell(row, sales_columns.unit_price)
        for col in (
            result_columns.bundle_price,
            result_columns.ratio,
            result_columns.split_unit_price,
            result_columns.split_amount,
        ):
            target = sheet.cell(row, col)
            if source.has_style:
                target._style = copy(source._style)
            target.number_format = "0.00"
        amount_source = sheet.cell(row, sales_columns.original_amount)
        for col in (
            result_columns.allocated_amount,
            result_columns.source_unit_price,
            result_columns.fee,
            result_columns.final_amount,
        ):
            if col is None:
                continue
            target = sheet.cell(row, col)
            if amount_source.has_style:
                target._style = copy(amount_source._style)
        for col in result_columns.all():
            sheet.cell(row, col).number_format = "0.00"

    header_fill = PatternFill("solid", fgColor="245C73")
    result_headers = {
        result_columns.bundle_price: SPLIT_RESULT_HEADERS[0],
        result_columns.ratio: SPLIT_RESULT_HEADERS[1],
        result_columns.split_unit_price: SPLIT_RESULT_HEADERS[2],
        result_columns.split_amount: SPLIT_RESULT_HEADERS[3],
        result_columns.allocated_amount: AMOUNT_RESULT_HEADERS[3],
        result_columns.source_unit_price: AMOUNT_RESULT_HEADERS[0],
        result_columns.fee: AMOUNT_RESULT_HEADERS[1],
        result_columns.final_amount: AMOUNT_RESULT_HEADERS[2],
    }
    for col, title in result_headers.items():
        if col is None:
            continue
        cell = sheet.cell(header_row, col, title)
        cell.fill = header_fill
        cell.font = Font(name="Microsoft YaHei UI", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _extend_auto_filter(sheet, header_row: int) -> None:
    end_column = get_column_letter(sheet.max_column)
    if sheet.auto_filter and sheet.auto_filter.ref:
        ref = sheet.auto_filter.ref
        if ":" in ref:
            start, end = ref.split(":", 1)
            end_row = "".join(ch for ch in end if ch.isdigit()) or str(sheet.max_row)
            sheet.auto_filter.ref = f"{start}:{end_column}{end_row}"
            return
    sheet.auto_filter.ref = f"A{header_row}:{end_column}{sheet.max_row}"


def _refresh_pivots_on_open(workbook) -> None:
    for sheet in workbook.worksheets:
        for pivot in getattr(sheet, "_pivots", ()):
            cache = getattr(pivot, "cache", None)
            if cache is not None:
                cache.enableRefresh = True
                cache.refreshOnLoad = True


def _freeze_cached_amounts(workbook, cached_workbook) -> int:
    frozen = 0
    for sheet in workbook.worksheets:
        if sheet.title not in cached_workbook.sheetnames:
            continue
        for header_row in range(1, min(sheet.max_row, 5) + 1):
            headers = _header_map(sheet, header_row)
            amount_columns = [
                headers[header]
                for header in ("订单金额", "应收合计")
                if header in headers
            ]
            if not amount_columns:
                continue
            cached_sheet = cached_workbook[sheet.title]
            for column in amount_columns:
                cached_values = cached_sheet.iter_rows(
                    min_row=header_row + 1,
                    max_row=sheet.max_row,
                    min_col=column,
                    max_col=column,
                    values_only=True,
                )
                for row, values in enumerate(cached_values, start=header_row + 1):
                    formula = sheet.cell(row, column).value
                    if isinstance(formula, str) and formula.startswith("="):
                        sheet.cell(row, column).value = values[0]
                        if values[0] is not None:
                            frozen += 1
            break
    return frozen


def _load_cached_order_amounts(cached_workbook) -> dict[str, float]:
    order_amounts: dict[str, float] = {}
    for sheet in cached_workbook.worksheets:
        for header_row in range(1, min(sheet.max_row, 5) + 1):
            headers = _header_map(sheet, header_row)
            if not {"网店订单号", "订单金额"}.issubset(headers):
                continue
            web_order_col = headers["网店订单号"]
            order_amount_col = headers["订单金额"]
            for row, values in enumerate(
                sheet.iter_rows(min_row=header_row + 1, values_only=True),
                start=header_row + 1,
            ):
                web_order = _text(values[web_order_col - 1])
                amount_value = values[order_amount_col - 1]
                if not web_order or _text(amount_value) == "":
                    continue
                amount = _number(amount_value, f"{sheet.title}第 {row} 行订单金额")
                if web_order in order_amounts and abs(order_amounts[web_order] - amount) > 1e-9:
                    raise SplitterError(f"Sheet1 网店订单 {web_order} 存在多个不同的订单金额。")
                order_amounts[web_order] = amount
            return order_amounts
    return order_amounts


def _prepare_result_sheet(
    workbook,
    source_sheet,
    header_row: int,
    charge_fee: bool,
) -> tuple[object, ResultLayout]:
    old_layout = _capture_column_layout(source_sheet)
    source_sales = _sales_columns(source_sheet, header_row)
    if "拆分结果" in workbook.sheetnames:
        workbook.remove(workbook["拆分结果"])
    result_sheet = workbook.copy_worksheet(source_sheet)
    result_sheet.title = "拆分结果"

    sales_positions = {
        name: getattr(source_sales, name)
        for name in SalesColumns.__dataclass_fields__
    }
    source_positions = {
        column: column for column in range(1, source_sheet.max_column + 1)
    }
    result_positions: dict[str, int] = {}
    amount_columns = (
        (
            ("source_unit_price", AMOUNT_RESULT_HEADERS[0]),
            ("fee", AMOUNT_RESULT_HEADERS[1]),
            ("final_amount", AMOUNT_RESULT_HEADERS[2]),
            ("allocated_amount", AMOUNT_RESULT_HEADERS[3]),
        )
        if charge_fee
        else (("allocated_amount", AMOUNT_RESULT_HEADERS[3]),)
    )
    insertions = (
        (
            source_sales.unit_price + 1,
            (
                ("bundle_price", SPLIT_RESULT_HEADERS[0]),
                ("ratio", SPLIT_RESULT_HEADERS[1]),
                ("split_unit_price", SPLIT_RESULT_HEADERS[2]),
                ("split_amount", SPLIT_RESULT_HEADERS[3]),
            ),
        ),
        (
            source_sales.original_amount + 1,
            amount_columns,
        ),
    )
    for insert_at, columns in sorted(insertions, reverse=True):
        count = len(columns)
        result_sheet.insert_cols(insert_at, amount=count)
        sales_positions = {
            name: (
                column + count
                if column is not None and column >= insert_at
                else column
            )
            for name, column in sales_positions.items()
        }
        source_positions = {
            old: column + count if column >= insert_at else column
            for old, column in source_positions.items()
        }
        result_positions = {
            name: column + count if column >= insert_at else column
            for name, column in result_positions.items()
        }
        for offset, (name, title) in enumerate(columns):
            column = insert_at + offset
            result_positions[name] = column
            result_sheet.cell(header_row, column, title)

    sales_columns = SalesColumns(**sales_positions)
    result_columns = ResultColumns(**result_positions)
    layout = ResultLayout(
        sales=sales_columns,
        result=result_columns,
        source_columns=tuple(source_positions.values()),
    )
    _copy_column_layout(result_sheet, old_layout, source_positions, result_columns)
    _style_new_columns(result_sheet, header_row, sales_columns, result_columns)
    _extend_auto_filter(result_sheet, header_row)
    return result_sheet, layout


def _zero_targets(sheet, row: int, result_columns: ResultColumns) -> None:
    for col in result_columns.all():
        sheet.cell(row, col, 0.0)


def _write_result_formulas(
    sheet,
    header_row: int,
    layout: ResultLayout,
) -> None:
    if (
        layout.result.source_unit_price is None
        or layout.result.fee is None
        or layout.result.final_amount is None
    ):
        return
    amount_letter = get_column_letter(layout.sales.original_amount)
    quantity_letter = get_column_letter(layout.sales.quantity)
    receivable_letter = get_column_letter(layout.sales.receivable)
    fee_letter = get_column_letter(layout.result.fee)
    for row in range(header_row + 1, sheet.max_row + 1):
        if all(
            sheet.cell(row, col).value is None
            for col in layout.source_columns
        ):
            continue
        sheet.cell(
            row,
            layout.result.source_unit_price,
            f"={amount_letter}{row}/{quantity_letter}{row}",
        )
        sheet.cell(
            row,
            layout.result.fee,
            f"=ROUND({amount_letter}{row}*1%,2)",
        )
        sheet.cell(
            row,
            layout.result.final_amount,
            f"={receivable_letter}{row}-{fee_letter}{row}",
        )


def _prepare_summary_sheet(
    workbook,
    result_sheet,
    header_row: int,
    layout: ResultLayout,
    web_targets: OrderedDict[str, float],
    charge_fee: bool,
):
    if "透视表" in workbook.sheetnames:
        workbook.remove(workbook["透视表"])
    summary = workbook.create_sheet("透视表")
    target_header = "平均值项:最终金额" if charge_fee else "应收合计"
    summary.append(["网店订单号", target_header, "求和项:摊后金额", "差异"])

    final_totals: OrderedDict[str, float] = (
        OrderedDict() if charge_fee else OrderedDict(web_targets)
    )
    allocated_totals: OrderedDict[str, float] = OrderedDict()
    for row in range(header_row + 1, result_sheet.max_row + 1):
        order = _text(result_sheet.cell(row, layout.sales.web_order_number).value)
        if not order:
            continue
        final_totals.setdefault(order, 0.0)
        allocated_totals.setdefault(order, 0.0)
        if charge_fee:
            try:
                original_amount = _number(
                    result_sheet.cell(row, layout.sales.original_amount).value or 0,
                    f"第 {row} 行原金额",
                )
                final_totals[order] = _round_two(
                    final_totals[order]
                    + original_amount
                    - _round_two(original_amount * 0.01)
                )
            except ValueError:
                pass
        allocated = result_sheet.cell(
            row, layout.result.allocated_amount
        ).value
        if _text(allocated) != "":
            try:
                allocated_totals[order] = _round_two(
                    allocated_totals[order]
                    + _number(allocated, f"第 {row} 行摊后金额")
                )
            except ValueError:
                pass

    for order in sorted(set(final_totals) | set(allocated_totals)):
        row = summary.max_row + 1
        target = _round_two(final_totals.get(order, 0.0))
        allocated = _round_two(allocated_totals.get(order, 0.0))
        difference = _round_two(allocated - target)
        summary.cell(row, 1, order)
        summary.cell(row, 2, target)
        summary.cell(row, 3, allocated)
        summary.cell(row, 4, difference)

    total_row = summary.max_row + 1
    target_total = _round_two(sum(final_totals.values()))
    allocated_total = _round_two(sum(allocated_totals.values()))
    summary.cell(total_row, 1, "总计")
    summary.cell(total_row, 2, target_total)
    summary.cell(total_row, 3, allocated_total)
    summary.cell(total_row, 4, _round_two(allocated_total - target_total))

    header_fill = PatternFill("solid", fgColor="245C73")
    for cell in summary[1]:
        cell.fill = header_fill
        cell.font = Font(name="Microsoft YaHei UI", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for cell in summary[total_row]:
        cell.font = Font(name="Microsoft YaHei UI", size=10, bold=True)
    for row in range(2, total_row + 1):
        summary.cell(row, 1).number_format = "@"
        for col in range(2, 5):
            summary.cell(row, col).number_format = "0.00_);[Red]\\(0.00\\)"
        if abs(float(summary.cell(row, 4).value or 0.0)) >= 0.005:
            summary.cell(row, 4).fill = PatternFill("solid", fgColor="FCE8E6")
            summary.cell(row, 4).font = Font(
                name="Microsoft YaHei UI",
                size=10,
                bold=True,
                color="C5221F",
            )
    summary.column_dimensions["A"].width = 36
    summary.column_dimensions["B"].width = 18.25
    summary.column_dimensions["C"].width = 13
    summary.column_dimensions["D"].width = 12.625
    summary.freeze_panes = "A2"
    if total_row > 2:
        summary.auto_filter.ref = f"A1:D{total_row - 1}"

    if "汇总" in workbook.sheetnames:
        workbook.remove(workbook["汇总"])
    totals = workbook.create_sheet("汇总")
    allocated_letter = get_column_letter(layout.result.allocated_amount)
    if charge_fee:
        original_amount_letter = get_column_letter(layout.sales.original_amount)
        fee_letter = get_column_letter(layout.result.fee)
        totals.append(
            ["最终金额", f"=SUM('拆分结果'!{original_amount_letter}:{original_amount_letter})"]
        )
        totals.append(
            ["摊后金额", f"=SUM('拆分结果'!{allocated_letter}:{allocated_letter})"]
        )
        totals.append(["手续费", f"=SUM('拆分结果'!{fee_letter}:{fee_letter})"])
        totals.append(["差异", "=B1-B2-B3"])
    else:
        totals.append(["应收合计", f"='透视表'!B{total_row}"])
        totals.append(
            ["摊后金额", f"=SUM('拆分结果'!{allocated_letter}:{allocated_letter})"]
        )
        totals.append(["差异", "=B1-B2"])
    totals.column_dimensions["A"].width = 14
    totals.column_dimensions["B"].width = 18
    for row in range(1, totals.max_row + 1):
        totals.cell(row, 1).font = Font(
            name="Microsoft YaHei UI",
            size=10,
            bold=True,
        )
        totals.cell(row, 2).number_format = "0.00_);[Red]\\(0.00\\)"
    return summary


def _enable_formula_recalculation(workbook) -> None:
    if workbook.calculation is None:
        workbook.calculation = CalcProperties()
    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True


def _mother_code(note: object) -> str:
    match = re.search(r"编号\s*[:：]\s*([^;；,，\s]+)", _text(note), re.IGNORECASE)
    return _code(match.group(1)) if match else ""


def _calculate_rows(
    sheet,
    header_row: int,
    layout: ResultLayout,
    ratio_data: RatioData,
    order_amounts: dict[str, float],
    charge_fee: bool = True,
    order_progress: OrderProgressCallback | None = None,
) -> tuple[SplitStats, OrderedDict[str, float]]:
    sales_columns = layout.sales
    result_columns = layout.result
    order_col = sales_columns.order_number
    logistics_col = sales_columns.logistics_number
    web_order_col = sales_columns.web_order_number
    total_col = sales_columns.receivable
    item_col = sales_columns.item_code
    quantity_col = sales_columns.quantity
    unit_price_col = sales_columns.unit_price
    note_col = sales_columns.note
    original_amount_col = sales_columns.original_amount

    base_groups: OrderedDict[tuple[str, str, int], list[int]] = OrderedDict()
    row_group_keys: dict[int, tuple[str, str, int]] = {}
    rows_by_web_order: OrderedDict[str, list[int]] = OrderedDict()
    split_web_orders: set[str] = set()
    active_rows: list[int] = []
    for row in range(header_row + 1, sheet.max_row + 1):
        if all(
            sheet.cell(row, col).value is None
            for col in layout.source_columns
        ):
            continue
        active_rows.append(row)
        logistics_number = _text(sheet.cell(row, logistics_col).value)
        order_number = _text(sheet.cell(row, order_col).value)
        web_order_number = _text(sheet.cell(row, web_order_col).value)
        if logistics_number:
            key = ("物流单号", logistics_number, 0)
        elif order_number:
            key = ("订单编号", order_number, 0)
        elif web_order_number:
            key = ("网店订单号", web_order_number, 0)
        else:
            key = ("单行", "", row)
        row_group_keys[row] = key
        base_groups.setdefault(key, []).append(row)
        if web_order_number:
            rows_by_web_order.setdefault(web_order_number, []).append(row)
            marker = (
                _text(sheet.cell(row, sales_columns.marker).value)
                if sales_columns.marker is not None
                else ""
            )
            if "拆分" in marker:
                split_web_orders.add(web_order_number)

    parent = {key: key for key in base_groups}

    def find(key: tuple[str, str, int]) -> tuple[str, str, int]:
        while parent[key] != key:
            parent[key] = parent[parent[key]]
            key = parent[key]
        return key

    def union(left: tuple[str, str, int], right: tuple[str, str, int]) -> None:
        left_root = find(left)
        right_root = find(right)
        if left_root != right_root:
            parent[right_root] = left_root

    for web_order in split_web_orders:
        keys = list(
            dict.fromkeys(
                row_group_keys[row] for row in rows_by_web_order[web_order]
            )
        )
        for key in keys[1:]:
            union(keys[0], key)

    groups: OrderedDict[tuple[str, str, int], list[int]] = OrderedDict()
    for row in active_rows:
        groups.setdefault(find(row_group_keys[row]), []).append(row)

    stats = SplitStats(orders=len(groups), rows=sum(len(rows) for rows in groups.values()))
    group_data: OrderedDict[tuple[str, str, int], dict[str, object]] = OrderedDict()
    web_groups: OrderedDict[str, list[tuple[str, str, int]]] = OrderedDict()

    for group_key, rows in groups.items():
        invalid_reason: str | None = None
        original_amounts: dict[int, float] = {}
        unit_prices: dict[int, float] = {}
        quantities: dict[int, float] = {}
        try:
            for row in rows:
                amount_value = sheet.cell(row, original_amount_col).value
                original_amounts[row] = (
                    0.0
                    if _text(amount_value) == ""
                    else _number(amount_value, f"第 {row} 行原金额")
                )
                price_value = sheet.cell(row, unit_price_col).value
                unit_prices[row] = (
                    0.0
                    if _text(price_value) == ""
                    else _number(price_value, f"第 {row} 行原单价")
                )
                quantity_value = sheet.cell(row, quantity_col).value
                quantities[row] = (
                    0.0
                    if _text(quantity_value) == ""
                    else _number(quantity_value, f"第 {row} 行数量")
                )
        except ValueError as exc:
            invalid_reason = str(exc)

        receivables: list[float] = []
        row_receivables: dict[int, float] = {}
        if invalid_reason is None:
            try:
                for row in rows:
                    value = sheet.cell(row, total_col).value
                    if _text(value) != "":
                        receivable = _number(value, f"第 {row} 行应收合计")
                        receivables.append(receivable)
                        row_receivables[row] = receivable
            except ValueError as exc:
                invalid_reason = str(exc)

        web_orders = tuple(
            dict.fromkeys(
                _text(sheet.cell(row, web_order_col).value)
                for row in rows
                if _text(sheet.cell(row, web_order_col).value)
            )
        )
        markers = [
            _text(sheet.cell(row, sales_columns.marker).value)
            for row in rows
        ] if sales_columns.marker is not None else []
        marker_group = any(
            "拆分" in marker or "合并" in marker
            for marker in markers
        )
        receivable_by_web: OrderedDict[str, float] = OrderedDict()
        if invalid_reason is None:
            if not receivables:
                invalid_reason = (
                    "应收合计为空或公式没有缓存值，请先用 Excel/WPS 打开并保存源文件后再上传"
                )
            elif marker_group and web_orders:
                for web_order in web_orders:
                    values = [
                        row_receivables[row]
                        for row in rows
                        if row in row_receivables
                        and _text(sheet.cell(row, web_order_col).value) == web_order
                    ]
                    if not values:
                        invalid_reason = f"网店订单号 {web_order} 的应收合计为空"
                        break
                    if any(abs(value - values[0]) > 1e-9 for value in values[1:]):
                        invalid_reason = f"网店订单号 {web_order} 存在多个不同的应收合计"
                        break
                    receivable_by_web[web_order] = _round_two(values[0])
            elif any(abs(value - receivables[0]) > 1e-9 for value in receivables[1:]):
                invalid_reason = f"同一{group_key[0]}存在多个不同的应收合计"

        group_receivable = (
            _round_two(sum(receivable_by_web.values()))
            if receivable_by_web
            else (receivables[0] if receivables else 0.0)
        )
        group_data[group_key] = {
            "rows": rows,
            "web_orders": web_orders,
            "original_amounts": original_amounts,
            "unit_prices": unit_prices,
            "quantities": quantities,
            "receivable": group_receivable,
            "receivable_by_web": receivable_by_web,
            "marker_group": marker_group,
            "fee": _round_two(
                sum(_round_two(amount * 0.01) for amount in original_amounts.values())
            ),
            "invalid_reason": invalid_reason,
        }
        for web_order in web_orders:
            web_groups.setdefault(web_order, []).append(group_key)

    group_gross_targets: dict[tuple[str, str, int], float] = {}
    group_targets: dict[tuple[str, str, int], float] = {}
    for group_key, data in group_data.items():
        if data["invalid_reason"] is None:
            receivable = _round_two(float(data["receivable"]))
            group_gross_targets[group_key] = receivable
            group_targets[group_key] = _round_two(
                receivable - float(data["fee"]) if charge_fee else receivable
            )

    external_group_targets: dict[tuple[str, str, int], float] = {}
    externally_covered_orders: dict[tuple[str, str, int], set[str]] = {}
    for web_order, external_total in order_amounts.items():
        valid_keys = [
            key
            for key in web_groups.get(web_order, [])
            if group_data[key]["invalid_reason"] is None
        ]
        if not valid_keys:
            continue
        weights = [abs(float(group_data[key]["receivable"])) for key in valid_keys]
        if sum(weights) < 1e-15:
            weights = [
                abs(sum(group_data[key]["original_amounts"].values()))
                for key in valid_keys
            ]
        total_weight = sum(weights)
        if total_weight < 1e-15:
            continue
        targets = _allocate_two(external_total, weights)
        for key, target in zip(valid_keys, targets):
            external_group_targets[key] = _round_two(
                external_group_targets.get(key, 0.0) + target
            )
            externally_covered_orders.setdefault(key, set()).add(web_order)

    for key, target in external_group_targets.items():
        if set(group_data[key]["web_orders"]).issubset(
            externally_covered_orders.get(key, set())
        ):
            group_gross_targets[key] = target
            group_targets[key] = _round_two(
                target - float(group_data[key]["fee"]) if charge_fee else target
            )

    total_orders = len(groups)
    for order_index, (group_key, rows) in enumerate(groups.items(), start=1):
        data = group_data[group_key]
        invalid_reason = data["invalid_reason"]
        display_order = group_key[1] or f"第 {group_key[2]} 行"
        if invalid_reason:
            stats.exceptional_orders += 1
            stats.warnings.append(f"订单 {display_order}：{invalid_reason}，拆分列已填 0。")
            for row in rows:
                _zero_targets(sheet, row, result_columns)
            if order_progress:
                order_progress(order_index, total_orders)
            continue

        original_amounts: dict[int, float] = data["original_amounts"]
        unit_prices: dict[int, float] = data["unit_prices"]
        quantities: dict[int, float] = data["quantities"]
        target_total = group_targets[group_key]

        if abs(target_total) < 1e-15:
            for row in rows:
                _zero_targets(sheet, row, result_columns)
            stats.unmatched_rows += len(rows)
            if order_progress:
                order_progress(order_index, total_orders)
            continue

        eligible_rows: list[int] = []
        for row in rows:
            if (
                abs(original_amounts[row]) >= 1e-15
                and abs(unit_prices[row]) >= 1e-15
            ):
                if quantities[row] == 0:
                    invalid_reason = f"第 {row} 行单价和金额不为 0，但数量为 0"
                    break
                eligible_rows.append(row)
            else:
                _zero_targets(sheet, row, result_columns)
                stats.unmatched_rows += 1

        if invalid_reason:
            stats.exceptional_orders += 1
            stats.warnings.append(f"订单 {display_order}：{invalid_reason}，拆分列已填 0。")
            for row in rows:
                _zero_targets(sheet, row, result_columns)
            if order_progress:
                order_progress(order_index, total_orders)
            continue

        recovered_zero_group = False
        if not eligible_rows:
            fallback_rows = [
                row for row in rows if abs(quantities[row]) >= 1e-15
            ]
            if fallback_rows:
                eligible_rows = fallback_rows
                recovered_zero_group = True
                stats.unmatched_rows -= len(fallback_rows)
            else:
                original_total = sum(original_amounts.values())
                if abs(original_total) < 1e-15:
                    invalid_reason = "整单单价或原金额合计为 0，且没有可分摊数量"
                else:
                    invalid_reason = "非零金额行的单价均为 0，无法按比例拆分"
        if invalid_reason:
            stats.exceptional_orders += 1
            stats.warnings.append(f"订单 {display_order}：{invalid_reason}，拆分列已填 0。")
            for row in rows:
                _zero_targets(sheet, row, result_columns)
            if order_progress:
                order_progress(order_index, total_orders)
            continue

        if not ratio_data.has_parent_ratios:
            row_weights: dict[int, float] = {}
            for row in eligible_rows:
                item_code = _code(sheet.cell(row, item_col).value)
                price = ratio_data.prices.get(item_code, 0.0)
                if price:
                    row_weights[row] = price
                    stats.matched_rows += 1
                else:
                    stats.unmatched_rows += 1
            if not row_weights:
                row_weights = {row: original_amounts[row] for row in eligible_rows}
            if abs(sum(row_weights.values())) < 1e-15:
                row_weights = {
                    row: abs(quantities[row]) for row in eligible_rows
                }
            weighted_rows = list(row_weights)
            allocated_amounts = _allocate_two(
                target_total,
                [row_weights[row] for row in weighted_rows],
            )
            row_amounts = dict(zip(weighted_rows, allocated_amounts))
            for row in eligible_rows:
                if row not in row_weights:
                    _zero_targets(sheet, row, result_columns)
                    continue
                ad = row_amounts[row]
                quantity = quantities[row]
                aa = ratio_data.prices.get(
                    _code(sheet.cell(row, item_col).value), 0.0
                )
                sheet.cell(row, result_columns.bundle_price, _round_two(aa))
                sheet.cell(
                    row,
                    result_columns.ratio,
                    _round_two(
                        ad / target_total / quantity if target_total else 0.0
                    ),
                )
                sheet.cell(
                    row,
                    result_columns.split_unit_price,
                    _round_two(ad / quantity),
                )
                sheet.cell(row, result_columns.split_amount, ad)
                sheet.cell(row, result_columns.allocated_amount, ad)
        else:
            allocation_groups: OrderedDict[str, list[int]] = OrderedDict()
            parents: dict[int, str] = {}
            for row in eligible_rows:
                parent = _mother_code(sheet.cell(row, note_col).value)
                parents[row] = parent
                key = f"PARENT:{parent}" if parent else f"ROW:{row}"
                allocation_groups.setdefault(key, []).append(row)

            group_weights = {
                key: sum(original_amounts[row] for row in group_rows)
                for key, group_rows in allocation_groups.items()
            }
            group_weight_total = sum(group_weights.values())
            if abs(group_weight_total) < 1e-15:
                group_weights = {
                    key: sum(
                        abs(unit_prices[row] * quantities[row])
                        for row in group_rows
                    )
                    for key, group_rows in allocation_groups.items()
                }
                group_weight_total = sum(group_weights.values())
            if abs(group_weight_total) < 1e-15:
                group_weights = {
                    key: sum(abs(quantities[row]) for row in group_rows)
                    for key, group_rows in allocation_groups.items()
                }
                group_weight_total = sum(group_weights.values())

            group_items = list(allocation_groups.items())
            group_targets_for_order = _allocate_two(
                target_total,
                [group_weights[key] for key, _ in group_items],
            )
            for (key, group_rows), group_target in zip(
                group_items,
                group_targets_for_order,
            ):
                parent = parents[group_rows[0]]
                entries = {
                    row: ratio_data.pairs.get(
                        (parent, _code(sheet.cell(row, item_col).value))
                    )
                    for row in group_rows
                }
                use_ratio = bool(parent) and all(
                    entry is not None
                    and abs(entry.ratio) >= 1e-15
                    and (
                        recovered_zero_group
                        or abs(entry.price) >= 1e-15
                    )
                    for entry in entries.values()
                )

                row_weights: dict[int, float] = {}
                if use_ratio:
                    child_rows: OrderedDict[str, list[int]] = OrderedDict()
                    for row in group_rows:
                        child_rows.setdefault(
                            _code(sheet.cell(row, item_col).value), []
                        ).append(row)
                    for child, duplicate_rows in child_rows.items():
                        entry = ratio_data.pairs[(parent, child)]
                        duplicate_weights = {
                            row: abs(unit_prices[row] * quantities[row])
                            for row in duplicate_rows
                        }
                        duplicate_total = sum(duplicate_weights.values())
                        if duplicate_total < 1e-15:
                            duplicate_weights = {row: 1.0 for row in duplicate_rows}
                            duplicate_total = float(len(duplicate_rows))
                        for row in duplicate_rows:
                            row_weights[row] = (
                                entry.ratio
                                * duplicate_weights[row]
                                / duplicate_total
                            )
                    if abs(sum(row_weights.values())) < 1e-15:
                        use_ratio = False

                if use_ratio:
                    stats.matched_rows += len(group_rows)
                else:
                    row_weights = {
                        row: abs(unit_prices[row] * quantities[row])
                        for row in group_rows
                    }
                    if abs(sum(row_weights.values())) < 1e-15:
                        row_weights = {
                            row: abs(original_amounts[row]) for row in group_rows
                        }
                    if abs(sum(row_weights.values())) < 1e-15:
                        row_weights = {
                            row: abs(quantities[row]) for row in group_rows
                        }
                    for row in group_rows:
                        entry = entries[row]
                        if (
                            entry is None
                            or abs(entry.price) < 1e-15
                            or abs(entry.ratio) < 1e-15
                        ):
                            stats.unmatched_rows += 1
                        else:
                            stats.matched_rows += 1

                row_amounts = _allocate_two(
                    group_target,
                    [row_weights[row] for row in group_rows],
                )
                for row, ad in zip(group_rows, row_amounts):
                    quantity = quantities[row]
                    aa = (
                        entries[row].price
                        if use_ratio and entries[row] is not None
                        else unit_prices[row]
                    )
                    sheet.cell(
                        row,
                        result_columns.bundle_price,
                        _round_two(aa),
                    )
                    sheet.cell(
                        row,
                        result_columns.ratio,
                        _round_two(
                            ad / target_total / quantity
                            if target_total
                            else 0.0
                        ),
                    )
                    sheet.cell(
                        row,
                        result_columns.split_unit_price,
                        _round_two(ad / quantity),
                    )
                    sheet.cell(row, result_columns.split_amount, ad)
                    sheet.cell(row, result_columns.allocated_amount, ad)

        ratio_weights = [
            abs(float(sheet.cell(row, result_columns.split_amount).value or 0.0))
            for row in rows
        ]
        if sum(ratio_weights) >= 1e-15:
            for row, ratio in zip(rows, _allocate_two(1.0, ratio_weights)):
                sheet.cell(row, result_columns.ratio, ratio)

        if order_progress:
            order_progress(order_index, total_orders)

    web_targets: OrderedDict[str, float] = OrderedDict()
    for group_key, data in group_data.items():
        if group_key not in group_gross_targets:
            continue
        web_orders = list(data["web_orders"])
        if not web_orders:
            continue
        gross_target = group_gross_targets[group_key]
        if data["marker_group"] and data["receivable_by_web"]:
            for order, amount in data["receivable_by_web"].items():
                web_targets[order] = _round_two(
                    web_targets.get(order, 0.0) + amount
                )
            continue
        if len(web_orders) == 1:
            web_targets[web_orders[0]] = _round_two(
                web_targets.get(web_orders[0], 0.0) + gross_target
            )
            continue

        rows_by_web: OrderedDict[str, list[int]] = OrderedDict(
            (order, []) for order in web_orders
        )
        for row in data["rows"]:
            order = _text(sheet.cell(row, web_order_col).value)
            if order in rows_by_web:
                rows_by_web[order].append(row)
        weights = [
            abs(
                sum(
                    float(
                        sheet.cell(row, result_columns.split_amount).value or 0.0
                    )
                    for row in rows_by_web[order]
                )
            )
            for order in web_orders
        ]
        if sum(weights) < 1e-15:
            weights = [
                sum(abs(data["original_amounts"][row]) for row in rows_by_web[order])
                for order in web_orders
            ]
        if sum(weights) < 1e-15:
            weights = [
                sum(abs(data["quantities"][row]) for row in rows_by_web[order])
                for order in web_orders
            ]
        if sum(weights) < 1e-15:
            weights = [1.0] * len(web_orders)
        for order, amount in zip(web_orders, _allocate_two(gross_target, weights)):
            web_targets[order] = _round_two(
                web_targets.get(order, 0.0) + amount
            )

    return stats, web_targets


def process_workbooks(
    ratio_path: str | Path,
    sales_path: str | Path,
    output_path: str | Path,
    progress: ProgressCallback | None = None,
    charge_fee: bool = True,
) -> SplitStats:
    """Generate a numeric result workbook that mirrors the reference formulas."""

    def report(percent: int, message: str) -> None:
        if progress:
            progress(percent, message)

    ratio_path = Path(ratio_path)
    sales_path = Path(sales_path)
    output_path = Path(output_path)
    if ratio_path.resolve() == sales_path.resolve():
        raise SplitterError("组合装占比表和销售单不能是同一个文件。")
    if sales_path.suffix.lower() != ".xlsx":
        raise SplitterError("销售单必须是 .xlsx 文件。")
    if output_path.suffix.lower() != ".xlsx":
        raise SplitterError("输出文件必须使用 .xlsx 扩展名。")
    if output_path.resolve() in {ratio_path.resolve(), sales_path.resolve()}:
        raise SplitterError("输出文件不能覆盖上传的源文件。")

    report(10, "读取组合装拆分占比表")
    ratio_data = _load_ratio_data(ratio_path)
    report(30, f"已载入 {len(ratio_data.prices)} 个子件")

    try:
        workbook = load_workbook(sales_path, data_only=False, keep_links=False)
        cached_workbook = load_workbook(
            sales_path, data_only=True, read_only=True, keep_links=False
        )
    except Exception as exc:
        raise SplitterError(f"无法打开销售单：{exc}") from exc

    try:
        source_sheet, header_row = _find_sheet(workbook, SALES_REQUIRED_HEADERS)
        if source_sheet is None:
            raise SplitterError("销售单中未找到包含完整销售字段的明细工作表。")
        _freeze_cached_amounts(workbook, cached_workbook)
        report(45, f"复制销售明细表“{source_sheet.title}”")
        result_sheet, layout = _prepare_result_sheet(
            workbook,
            source_sheet,
            header_row,
            charge_fee,
        )
        order_amounts = (
            _load_cached_order_amounts(cached_workbook) if charge_fee else {}
        )
        mode_label = "扣除手续费后" if charge_fee else "按应收合计"
        report(60, f"{mode_label}计算拆分单价")
        stats, web_targets = _calculate_rows(
            result_sheet,
            header_row,
            layout,
            ratio_data,
            order_amounts,
            charge_fee=charge_fee,
            order_progress=lambda done, total: report(
                60 + int((done / total) * 28) if total else 88,
                f"正在拆分订单 {done}/{total}",
            ),
        )
        _write_result_formulas(result_sheet, header_row, layout)
        _prepare_summary_sheet(
            workbook,
            result_sheet,
            header_row,
            layout,
            web_targets,
            charge_fee,
        )
        _enable_formula_recalculation(workbook)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        report(90, "保存结果工作簿")
        _refresh_pivots_on_open(workbook)
        workbook.save(output_path)
        stats.output_path = str(output_path)
        report(100, "拆分完成")
        return stats
    except SplitterError:
        raise
    except Exception as exc:
        raise SplitterError(f"生成结果时发生错误：{exc}") from exc
    finally:
        workbook.close()
        cached_workbook.close()


def default_output_path(sales_path: str | Path) -> Path:
    from datetime import datetime

    source = Path(sales_path)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return source.with_name(f"{source.stem}_已拆单价_{stamp}.xlsx")
