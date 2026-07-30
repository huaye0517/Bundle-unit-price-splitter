from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from splitter import (
    SplitterError,
    _enable_formula_recalculation,
    _freeze_cached_amounts,
    _round_two,
    process_workbooks,
    ratio_file_info,
    update_ratio_data,
)


SALES_HEADERS = [
    "标记", "订单编号", "订单状态", "结算状态", "销售渠道", "处理时间", "付款时间", "发货仓库",
    "物流公司", "物流单号", "网店订单号", "发货时间", "订单类型", "应收合计", "货品数量", "货品摘要",
    "客户账号", "收货人", "手机", "收货地址", "合并备注", "货品编号", "货品名称", "规格", "数量", "单价",
    "优惠", "折扣", "金额", "锁定待发", "备注", "达人ID", "达人名称", "赠品",
]


def make_ratio(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet2"
    ws.append(["组合装编号", "单品编号", "金额", "分摊比例", "单价", "数量"])
    ws.append(["A", "SKU1", 100, 0.4, 40, 1])
    ws.append(["B", "SKU1", 200, 0.5, 100, 1])  # must be ignored: VLOOKUP first match
    ws.append(["A", "SKU2", 100, 0.6, 60, 2])
    wb.save(path)


def make_new_ratio(path: Path) -> None:
    wb = Workbook()
    wb.active.title = "sheet"
    ws = wb.create_sheet("sheet1")
    headers = ["母件编号", "编号", "名称", "规格", "条码", "长(cm)", "宽(cm)", "高(cm)", "重量(g)", "体积(cm³)", "单位", "数量", "执行价格", "分摊金额", "分摊比例", "上传SN", "赠品", "换算率"]
    ws.append(headers)
    ws.append(["A", "SKU1", "商品1", "", "", "", "", "", "", "", "件", 2, 40, 80, 0.4, "否", "否", ""])
    ws.append(headers)  # merged exports can contain repeated header rows
    ws.append(["A", "SKU2", "商品2", "", "", "", "", "", "", "", "件", 2, 60, 120, 0.6, "否", "否", ""])
    wb.save(path)


def make_sales(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "销售明细"
    ws.append(SALES_HEADERS)
    base = [""] * len(SALES_HEADERS)
    for index, (code, qty) in enumerate([("MISS", 1), ("SKU1", 1), ("SKU2", 2)]):
        row = base.copy()
        row[1] = "ORDER-1"
        row[10] = "WEB-1"
        row[13] = 200 if index == 0 else None
        row[21] = code
        row[24] = qty
        row[25] = 10
        row[28] = qty * 10
        ws.append(row)
    wb.create_sheet("其他工作表")["A1"] = "保留"
    wb.save(path)


def add_order_amounts(path: Path, amounts: list[tuple[str, float]]) -> None:
    wb = load_workbook(path)
    if "Sheet1" in wb.sheetnames:
        del wb["Sheet1"]
    ws = wb.create_sheet("Sheet1")
    headers = [""] * 14
    headers[11] = "网店订单号"
    headers[13] = "订单金额"
    ws.append(headers)
    for web_order, amount in amounts:
        row = [""] * 14
        row[11] = web_order
        row[13] = amount
        ws.append(row)
    wb.save(path)
    wb.close()


class SplitterTests(unittest.TestCase):
    def test_round_two_uses_standard_half_up_rounding(self):
        self.assertEqual(_round_two(1.005), 1.01)
        self.assertEqual(_round_two(-1.005), -1.01)
        self.assertEqual(_round_two(1.004), 1.0)

    def test_rounding_remainder_keeps_order_total_at_two_decimals(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            ratio, sales, output = (
                root / "ratio.xlsx",
                root / "sales.xlsx",
                root / "output.xlsx",
            )
            make_ratio(ratio)
            wb = Workbook()
            ws = wb.active
            ws.title = "销售明细"
            ws.append(SALES_HEADERS)
            for index in range(3):
                row = [""] * len(SALES_HEADERS)
                row[1] = f"ORDER-{index + 1}"
                row[9] = "TRACKING-1"
                row[10] = "WEB-1"
                row[13] = 100.01
                row[21] = "SKU1"
                row[24] = 1
                row[25] = 1
                row[28] = 1
                ws.append(row)
            wb.save(sales)
            wb.close()

            process_workbooks(ratio, sales, output)

            result_book = load_workbook(output, data_only=False)
            result = result_book["拆分结果"]
            self.assertEqual(
                [result.cell(row, 30).value for row in range(2, 5)],
                [33.33, 33.33, 33.32],
            )
            self.assertAlmostEqual(
                sum(result.cell(row, 30).value for row in range(2, 5)),
                99.98,
                places=8,
            )
            for row in range(2, 5):
                for column in (27, 28, 29, 30, 37):
                    value = result.cell(row, column).value
                    self.assertEqual(value, _round_two(value))
                    self.assertEqual(result.cell(row, column).number_format, "0.00")
                self.assertEqual(
                    result.cell(row, 30).value,
                    result.cell(row, 37).value,
                )
            summary = result_book["透视表"]
            self.assertEqual(summary["B2"].value, 2.97)
            self.assertEqual(summary["C2"].value, 99.98)
            self.assertEqual(summary["D2"].value, 97.01)
            result_book.close()

    def test_missing_calculation_properties_are_initialized(self):
        workbook = Workbook()
        workbook.calculation = None

        _enable_formula_recalculation(workbook)

        self.assertEqual(workbook.calculation.calcMode, "auto")
        self.assertTrue(workbook.calculation.fullCalcOnLoad)
        self.assertTrue(workbook.calculation.forceFullCalc)
        workbook.close()

    def test_cached_receivable_formulas_are_frozen_as_values(self):
        formula_book = Workbook()
        formula_sheet = formula_book.active
        formula_sheet.title = "销售明细"
        formula_sheet.append(SALES_HEADERS)
        formula_sheet.append([None] * 10 + ["WEB-1", None, None, "=SUMIF(...)" ] + [None] * 20)
        formula_sheet.append([None] * 10 + ["WEB-1", None, None, "=SUMIF(...)" ] + [None] * 20)

        cached_book = Workbook()
        cached_sheet = cached_book.active
        cached_sheet.title = "销售明细"
        cached_sheet.append(SALES_HEADERS)
        cached_sheet.append([None] * 10 + ["WEB-1", None, None, 187.5] + [None] * 20)
        cached_sheet.append([None] * 10 + ["WEB-1", None, None, None] + [None] * 20)

        frozen = _freeze_cached_amounts(formula_book, cached_book)

        self.assertEqual(frozen, 1)
        self.assertEqual(formula_sheet["N2"].value, 187.5)
        self.assertIsNone(formula_sheet["N3"].value)
        formula_book.close()
        cached_book.close()

    def test_repeated_receivable_total_is_counted_once_and_fee_is_deducted(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            ratio, sales, output = root / "ratio.xlsx", root / "sales.xlsx", root / "output.xlsx"
            make_ratio(ratio)
            wb = Workbook()
            ws = wb.active
            ws.title = "销售明细"
            ws.append(SALES_HEADERS)
            for index in range(10):
                row = [""] * len(SALES_HEADERS)
                row[1] = f"ORDER-{index + 1}"
                row[10] = "PO-260629-323800286823316"
                row[13] = 187.5
                row[21] = "SKU1"
                row[24] = 1
                row[25] = 18.75
                row[28] = 18.75
                ws.append(row)
            wb.save(sales)
            wb.close()

            stats = process_workbooks(ratio, sales, output)

            self.assertEqual(stats.orders, 10)
            self.assertEqual(stats.exceptional_orders, 0)
            result_book = load_workbook(output, data_only=False)
            result = result_book["拆分结果"]
            self.assertEqual([result.cell(row, 14).value for row in range(2, 12)], [187.5] * 10)
            net_total = 10 * 187.31
            self.assertAlmostEqual(sum(result.cell(row, 30).value for row in range(2, 12)), net_total, places=8)
            self.assertAlmostEqual(sum(result.cell(row, 37).value for row in range(2, 12)), net_total, places=8)
            for row in range(2, 12):
                self.assertAlmostEqual(result.cell(row, 30).value, result.cell(row, 37).value, places=8)
                self.assertEqual(result.cell(row, 34).value, f"=AG{row}/Y{row}")
                self.assertEqual(result.cell(row, 35).value, f"=ROUND(AG{row}*1%,2)")
                self.assertEqual(result.cell(row, 36).value, f"=N{row}-AI{row}")
            self.assertEqual(result["AH1"].value, "单价2")
            self.assertEqual(result["AI1"].value, "手续费1%")
            self.assertEqual(result["AJ1"].value, "最终金额")
            self.assertEqual(result["AK1"].value, "摊后金额")
            self.assertEqual(result["AH2"].value, "=AG2/Y2")
            self.assertEqual(result["AI2"].value, "=ROUND(AG2*1%,2)")
            self.assertEqual(result["AJ2"].value, "=N2-AI2")
            self.assertEqual(result["AP1"].value, "赠品")
            self.assertEqual(result.auto_filter.ref, "A1:AP11")
            summary = result_book["透视表"]
            self.assertEqual(
                [summary.cell(1, col).value for col in range(1, 5)],
                ["网店订单号", "平均值项:最终金额", "求和项:摊后金额", "差异"],
            )
            self.assertEqual(summary["A2"].value, "PO-260629-323800286823316")
            self.assertEqual(summary["B2"].value, 185.6)
            self.assertAlmostEqual(summary["C2"].value, net_total, places=8)
            self.assertEqual(summary["D2"].value, 1687.5)
            self.assertEqual(
                summary["D2"].number_format,
                "0.00_);[Red]\\(0.00\\)",
            )
            self.assertEqual(summary["A3"].value, "总计")
            self.assertEqual(summary["B3"].value, 185.6)
            self.assertEqual(summary["C3"].value, net_total)
            totals = result_book["汇总"]
            self.assertEqual(totals["A1"].value, "最终金额")
            self.assertEqual(totals["B1"].value, "=SUM('拆分结果'!AG:AG)")
            self.assertEqual(totals["A2"].value, "摊后金额")
            self.assertEqual(totals["B2"].value, "=SUM('拆分结果'!AK:AK)")
            self.assertEqual(totals["A3"].value, "手续费")
            self.assertEqual(totals["B3"].value, "=SUM('拆分结果'!AI:AI)")
            self.assertEqual(totals["A4"].value, "差异")
            self.assertEqual(totals["B4"].value, "=B1-B2-B3")
            self.assertEqual(summary["D3"].value, 1687.5)
            self.assertEqual(len(getattr(summary, "_pivots", [])), 0)
            self.assertEqual(result_book.calculation.calcMode, "auto")
            self.assertTrue(result_book.calculation.fullCalcOnLoad)
            self.assertTrue(result_book.calculation.forceFullCalc)
            result_book.close()

    def test_summary_uses_gross_receivable_and_fee_net_split(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            ratio, sales, output = root / "ratio.xlsx", root / "sales.xlsx", root / "output.xlsx"
            make_ratio(ratio)
            wb = Workbook()
            ws = wb.active
            ws.title = "销售明细"
            ws.append(SALES_HEADERS)
            for order, gross, original_amount in (
                ("260531-420353175082566", 24.0, 14.75),
                ("260531-420353175082566", 24.0, 14.75),
                ("260531-421947010581809", 73.5, 43.065),
                ("260531-421947010581809", 73.5, 43.065),
            ):
                row = [""] * len(SALES_HEADERS)
                row[1] = order
                row[10] = order
                row[13] = gross
                row[21] = "SKU1"
                row[24] = 1
                row[25] = original_amount
                row[28] = original_amount
                ws.append(row)
            wb.save(sales)
            wb.close()

            process_workbooks(ratio, sales, output)

            result = load_workbook(output, data_only=False)
            summary = result["透视表"]
            self.assertEqual(summary["B2"].value, 29.2)
            self.assertAlmostEqual(summary["C2"].value, 23.7, places=8)
            self.assertEqual(summary["B3"].value, 85.27)
            self.assertEqual(summary["C3"].value, 72.64)
            self.assertEqual(summary["D2"].value, -5.5)
            self.assertEqual(summary["D3"].value, -12.63)
            result.close()

    def test_conflicting_receivable_totals_mark_only_that_order_exceptional(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            ratio, sales, output = root / "ratio.xlsx", root / "sales.xlsx", root / "output.xlsx"
            make_ratio(ratio)
            make_sales(sales)
            wb = load_workbook(sales)
            ws = wb["销售明细"]
            ws["N2"] = 187.5
            ws["N3"] = 188
            ws["N4"] = 187.5
            wb.save(sales)
            wb.close()

            stats = process_workbooks(ratio, sales, output)

            self.assertEqual(stats.exceptional_orders, 1)
            self.assertTrue(any("多个不同的应收合计" in warning for warning in stats.warnings))
            result = load_workbook(output, data_only=True)["拆分结果"]
            for row in range(2, 5):
                self.assertEqual(result.cell(row, 30).value, 0)
                self.assertEqual(result.cell(row, 37).value, 0)

    def test_blank_web_order_falls_back_to_internal_order_number(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            ratio, sales, output = root / "ratio.xlsx", root / "sales.xlsx", root / "output.xlsx"
            make_ratio(ratio)
            make_sales(sales)
            wb = load_workbook(sales)
            ws = wb["销售明细"]
            for row in range(2, 5):
                ws.cell(row, 11, None)
                ws.cell(row, 14, 187.5)
            wb.save(sales)
            wb.close()

            stats = process_workbooks(ratio, sales, output)

            self.assertEqual(stats.orders, 1)
            self.assertEqual(stats.exceptional_orders, 0)
            result = load_workbook(output, data_only=True)["拆分结果"]
            self.assertAlmostEqual(sum(result.cell(row, 30).value for row in range(2, 5)), 187.1, places=8)
            self.assertAlmostEqual(sum(result.cell(row, 37).value for row in range(2, 5)), 187.1, places=8)

    def test_append_ratio_data_adds_new_bundle_items_without_overwriting(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            current, incoming = root / "current.xlsx", root / "incoming.xlsx"
            make_new_ratio(current)
            wb = Workbook()
            ws = wb.active
            ws.title = "总表"
            headers = ["母件编号", "编号", "名称", "规格", "条码", "长(cm)", "宽(cm)", "高(cm)", "重量(g)", "体积(cm³)", "单位", "数量", "执行价格", "分摊金额", "分摊比例", "上传SN", "赠品", "换算率"]
            ws.append(headers)
            ws.append(["A", "SKU1", "重复商品", "", "", "", "", "", "", "", "件", 1, 999, 999, 1, "否", "否", ""])
            ws.append(["B", "SKU3", "新增商品", "", "", "", "", "", "", "", "件", 1, 30, 30, 1, "否", "否", ""])
            wb.save(incoming)
            wb.close()

            stats = update_ratio_data(current, incoming, "append")

            self.assertEqual(stats.added_rows, 1)
            self.assertEqual(stats.source_rows, 3)
            self.assertEqual(stats.unique_items, 3)
            unique_items, source_rows = ratio_file_info(current)
            self.assertEqual((unique_items, source_rows), (3, 3))
            result = load_workbook(current, data_only=True)
            data_sheet = max(result.worksheets, key=lambda sheet: sheet.max_row)
            rows = list(data_sheet.iter_rows(min_row=2, values_only=True))
            self.assertEqual(sum(1 for row in rows if row[0] == "A" and row[1] == "SKU1"), 1)
            self.assertTrue(any(row[0] == "B" and row[1] == "SKU3" and row[12] == 30 for row in rows))
            result.close()

    def test_replace_ratio_data_uses_incoming_workbook(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            current, incoming = root / "current.xlsx", root / "incoming.xlsx"
            make_new_ratio(current)
            make_ratio(incoming)

            stats = update_ratio_data(current, incoming, "replace")

            self.assertEqual(stats.added_rows, 3)
            self.assertEqual(stats.source_rows, 3)
            self.assertEqual(stats.unique_items, 2)
            result = load_workbook(current, data_only=True)
            self.assertIn("Sheet2", result.sheetnames)
            self.assertNotIn("sheet1", result.sheetnames)
            result.close()

    def test_supports_new_ratio_export_format(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            ratio, sales, output = root / "ratio.xlsx", root / "sales.xlsx", root / "output.xlsx"
            make_new_ratio(ratio)
            make_sales(sales)
            sales_book = load_workbook(sales)
            sales_sheet = sales_book["销售明细"]
            sales_sheet["AE3"] = "编号:A;"
            sales_sheet["AE4"] = "编号：a；"
            sales_book.save(sales)
            sales_book.close()

            stats = process_workbooks(ratio, sales, output)

            self.assertEqual(stats.matched_rows, 2)
            result = load_workbook(output, data_only=True)["拆分结果"]
            self.assertEqual(result["AA3"].value, 40)
            self.assertEqual(result["AA4"].value, 60)
            self.assertAlmostEqual(
                sum(result.cell(row, 37).value for row in range(2, 5)),
                199.6,
                places=12,
            )

    def test_formula_chain_and_first_match(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            ratio, sales, output = root / "ratio.xlsx", root / "sales.xlsx", root / "output.xlsx"
            make_ratio(ratio)
            make_sales(sales)
            progress_events = []
            stats = process_workbooks(ratio, sales, output, progress=lambda value, message: progress_events.append((value, message)))

            self.assertEqual(stats.orders, 1)
            self.assertEqual(stats.matched_rows, 2)
            self.assertEqual(stats.unmatched_rows, 1)
            self.assertEqual(progress_events[-1][0], 100)
            self.assertTrue(any("正在拆分订单 1/1" in message for _, message in progress_events))
            wb = load_workbook(output, data_only=True)
            ws = wb["拆分结果"]
            self.assertEqual(ws["AA2"].value, 0)
            self.assertEqual(ws["AD2"].value, 0)
            self.assertEqual(ws["AA3"].value, 40)  # first SKU1, not 100
            self.assertEqual(ws["AA3"].number_format, "0.00")
            self.assertEqual(ws["AB3"].number_format, "0.00")
            self.assertEqual(ws["AK3"].number_format, "0.00")
            self.assertAlmostEqual(ws["AB3"].value, 40 / 100 / 1, places=14)
            self.assertAlmostEqual(ws["AC3"].value, 79.84, places=12)
            self.assertAlmostEqual(ws["AD3"].value, 79.84, places=12)
            self.assertAlmostEqual(ws["AA4"].value, 60, places=12)
            self.assertAlmostEqual(ws["AC4"].value, 59.88, places=12)
            self.assertAlmostEqual(ws["AD4"].value, 119.76, places=12)
            self.assertAlmostEqual(ws["AD3"].value + ws["AD4"].value, 199.6, places=12)
            self.assertAlmostEqual(ws["AK3"].value + ws["AK4"].value, 199.6, places=12)
            self.assertIn("销售明细", wb.sheetnames)
            self.assertIn("其他工作表", wb.sheetnames)
            self.assertEqual(wb["其他工作表"]["A1"].value, "保留")
            wb.close()

    def test_all_unmatched_order_allocates_receivable_by_original_amount(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            ratio, sales, output = root / "ratio.xlsx", root / "sales.xlsx", root / "output.xlsx"
            make_ratio(ratio)
            make_sales(sales)
            wb = load_workbook(sales)
            ws = wb["销售明细"]
            ws["V3"] = "MISS2"
            ws["V4"] = "MISS3"
            wb.save(sales)
            wb.close()
            stats = process_workbooks(ratio, sales, output)
            self.assertEqual(stats.exceptional_orders, 0)
            self.assertFalse(stats.warnings)
            result = load_workbook(output, data_only=True)["拆分结果"]
            for row in range(2, 5):
                self.assertEqual(result.cell(row, 27).value, 0)
                self.assertEqual(result.cell(row, 29).value, 49.9)
            self.assertEqual(result.cell(2, 30).value, 49.9)
            self.assertEqual(result.cell(3, 30).value, 49.9)
            self.assertEqual(result.cell(4, 30).value, 99.8)
            self.assertEqual(result.cell(2, 37).value, 49.9)
            self.assertEqual(result.cell(3, 37).value, 49.9)
            self.assertEqual(result.cell(4, 37).value, 99.8)
            self.assertEqual(sum(result.cell(row, 37).value for row in range(2, 5)), 199.6)

    def test_internal_orders_are_balanced_independently(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            ratio, sales, output = root / "ratio.xlsx", root / "sales.xlsx", root / "output.xlsx"
            make_ratio(ratio)
            make_sales(sales)
            wb = load_workbook(sales)
            ws = wb["销售明细"]
            ws["K2"] = "WEB-A"
            ws["K3"] = "WEB-A"
            ws["K4"] = "WEB-B"
            ws["B4"] = "ORDER-2"
            ws["N2"] = 50
            ws["N3"] = None
            ws["N4"] = 80
            wb.save(sales)
            wb.close()

            stats = process_workbooks(ratio, sales, output)
            self.assertEqual(stats.orders, 2)
            result = load_workbook(output, data_only=True)["拆分结果"]
            self.assertAlmostEqual(result["AD2"].value + result["AD3"].value, 49.8, places=12)
            self.assertAlmostEqual(result["AD4"].value, 79.8, places=12)
            self.assertAlmostEqual(result["AK2"].value + result["AK3"].value, 49.8, places=12)
            self.assertAlmostEqual(result["AK4"].value, 79.8, places=12)

    def test_sheet1_order_amounts_are_allocated_by_web_order(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            ratio, sales, output = root / "ratio.xlsx", root / "sales.xlsx", root / "output.xlsx"
            make_ratio(ratio)
            make_sales(sales)
            wb = load_workbook(sales)
            ws = wb["销售明细"]
            ws["K2"] = "WEB-A"
            ws["K3"] = "WEB-A"
            ws["K4"] = "WEB-B"
            ws["N2"] = 130
            ws["N3"] = 130
            ws["N4"] = 130
            wb.save(sales)
            wb.close()
            add_order_amounts(sales, [("WEB-A", 50), ("WEB-B", 80)])

            stats = process_workbooks(ratio, sales, output)

            self.assertEqual(stats.orders, 1)
            self.assertEqual(stats.exceptional_orders, 0)
            result = load_workbook(output, data_only=True)["拆分结果"]
            self.assertAlmostEqual(sum(result.cell(row, 30).value for row in range(2, 5)), 129.6, places=12)
            self.assertAlmostEqual(sum(result.cell(row, 37).value for row in range(2, 5)), 129.6, places=12)

    def test_zero_original_amount_does_not_participate_in_allocation(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            ratio, sales, output = root / "ratio.xlsx", root / "sales.xlsx", root / "output.xlsx"
            make_ratio(ratio)
            make_sales(sales)
            wb = load_workbook(sales)
            ws = wb["销售明细"]
            ws["AC3"] = 0
            wb.save(sales)
            wb.close()

            process_workbooks(ratio, sales, output)

            result = load_workbook(output, data_only=True)["拆分结果"]
            self.assertEqual(result["AA3"].value, 0)
            self.assertEqual(result["AD3"].value, 0)
            self.assertEqual(result["AK3"].value, 0)
            self.assertAlmostEqual(sum(result.cell(row, 30).value for row in range(2, 5)), 199.7, places=12)
            self.assertAlmostEqual(sum(result.cell(row, 37).value for row in range(2, 5)), 199.7, places=12)

    def test_sheet1_web_order_amount_overrides_repeated_source_totals(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            ratio, sales, output = root / "ratio.xlsx", root / "sales.xlsx", root / "output.xlsx"
            make_ratio(ratio)
            make_sales(sales)
            wb = load_workbook(sales)
            ws = wb["销售明细"]
            ws["K2"] = "WEB-SHARED"
            ws["K3"] = "WEB-SHARED"
            ws["K4"] = "WEB-SHARED"
            ws["B4"] = "ORDER-2"
            ws["N2"] = 50
            ws["N3"] = None
            ws["N4"] = 80
            wb.save(sales)
            wb.close()
            add_order_amounts(sales, [("WEB-SHARED", 130)])

            stats = process_workbooks(ratio, sales, output)

            self.assertEqual(stats.orders, 2)
            self.assertEqual(stats.exceptional_orders, 0)
            result = load_workbook(output, data_only=True)["拆分结果"]
            self.assertAlmostEqual(sum(result.cell(row, 30).value for row in range(2, 5)), 129.6, places=12)
            self.assertAlmostEqual(sum(result.cell(row, 37).value for row in range(2, 5)), 129.6, places=12)

    def test_uses_sales_sheet_with_most_rows(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            ratio, sales, output = root / "ratio.xlsx", root / "sales.xlsx", root / "output.xlsx"
            make_ratio(ratio)
            make_sales(sales)
            wb = load_workbook(sales)
            source = wb["销售明细"]
            corrected = wb.copy_worksheet(source)
            corrected.title = "销售明细修正版"
            row = [""] * len(SALES_HEADERS)
            row[1] = "ORDER-2"
            row[10] = "WEB-2"
            row[13] = 30
            row[21] = "SKU1"
            row[24] = 1
            row[25] = 30
            row[28] = 30
            corrected.append(row)
            wb.save(sales)
            wb.close()

            stats = process_workbooks(ratio, sales, output)
            self.assertEqual(stats.orders, 2)
            result = load_workbook(output, data_only=True)["拆分结果"]
            self.assertEqual(result.max_row, 5)
            self.assertEqual(result["K5"].value, "WEB-2")
            self.assertAlmostEqual(result["AD5"].value, 29.7, places=12)

    def test_same_web_order_splits_each_internal_order_and_sums_summary(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            ratio, sales, output = root / "ratio.xlsx", root / "sales.xlsx", root / "output.xlsx"
            make_new_ratio(ratio)
            wb = Workbook()
            ws = wb.active
            ws.title = "销售明细"
            ws.append(SALES_HEADERS)
            for order, total, values in (
                ("ORDER-A", 100, (("SKU1", 40), ("SKU2", 60))),
                ("ORDER-B", 50, (("SKU1", 20), ("SKU2", 30))),
            ):
                for code, amount in values:
                    row = [""] * len(SALES_HEADERS)
                    row[1] = order
                    row[10] = "WEB-SHARED"
                    row[13] = total
                    row[21] = code
                    row[24] = 1
                    row[25] = amount
                    row[28] = amount
                    row[30] = "原组合装名称:测试,编号:a;"
                    ws.append(row)
            wb.save(sales)
            wb.close()

            stats = process_workbooks(ratio, sales, output)

            self.assertEqual(stats.orders, 2)
            self.assertEqual(stats.exceptional_orders, 0)
            result_book = load_workbook(output, data_only=True)
            result = result_book["拆分结果"]
            self.assertEqual([result.cell(row, 27).value for row in range(2, 6)], [40, 60, 40, 60])
            self.assertAlmostEqual(
                sum(result.cell(row, 30).value for row in range(2, 4)),
                99,
                places=12,
            )
            self.assertAlmostEqual(
                sum(result.cell(row, 30).value for row in range(4, 6)),
                49.5,
                places=12,
            )
            summary = result_book["透视表"]
            self.assertEqual(summary["A2"].value, "WEB-SHARED")
            self.assertAlmostEqual(summary["B2"].value, 148.5, places=12)
            self.assertAlmostEqual(summary["C2"].value, 148.5, places=12)
            result_book.close()

    def test_same_logistics_number_uses_receivable_once_across_web_orders(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            ratio, sales, output = root / "ratio.xlsx", root / "sales.xlsx", root / "output.xlsx"
            make_new_ratio(ratio)
            wb = Workbook()
            ws = wb.active
            ws.title = "销售明细"
            ws.append(SALES_HEADERS)
            for web_order, code, amount in (
                ("WEB-A", "SKU1", 40),
                ("WEB-B", "SKU2", 60),
            ):
                row = [""] * len(SALES_HEADERS)
                row[1] = "ORDER-SHARED"
                row[9] = "TRACKING-SHARED"
                row[10] = web_order
                row[13] = 100
                row[21] = code
                row[24] = 1
                row[25] = amount
                row[28] = amount
                row[30] = "原组合装名称:测试,编号:a;"
                ws.append(row)
            wb.save(sales)
            wb.close()

            stats = process_workbooks(ratio, sales, output)

            self.assertEqual(stats.orders, 1)
            self.assertEqual(stats.exceptional_orders, 0)
            result_book = load_workbook(output, data_only=True)
            result = result_book["拆分结果"]
            self.assertAlmostEqual(result["AD2"].value, 39.6, places=12)
            self.assertAlmostEqual(result["AD3"].value, 59.4, places=12)
            self.assertLessEqual(result["AC2"].value, result["Z2"].value)
            self.assertLessEqual(result["AC3"].value, result["Z3"].value)
            summary = result_book["透视表"]
            self.assertEqual(summary["A2"].value, "WEB-A")
            self.assertAlmostEqual(summary["B2"].value, 39.6, places=12)
            self.assertAlmostEqual(summary["C2"].value, 39.6, places=12)
            self.assertEqual(summary["A3"].value, "WEB-B")
            self.assertAlmostEqual(summary["B3"].value, 59.4, places=12)
            self.assertAlmostEqual(summary["C3"].value, 59.4, places=12)
            result_book.close()

    def test_partial_parent_ratio_normalizes_and_missing_pair_uses_z(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            ratio, sales, output = root / "ratio.xlsx", root / "sales.xlsx", root / "output.xlsx"
            make_new_ratio(ratio)
            wb = Workbook()
            ws = wb.active
            ws.title = "销售明细"
            ws.append(SALES_HEADERS)
            for order, code, price, amount in (
                ("ORDER-PARTIAL", "SKU1", 100, 100),
                ("ORDER-FALLBACK", "SKU1", 30, 30),
                ("ORDER-FALLBACK", "MISS", 70, 70),
            ):
                row = [""] * len(SALES_HEADERS)
                row[1] = order
                row[10] = f"WEB-{order}"
                row[13] = 100
                row[21] = code
                row[24] = 1
                row[25] = price
                row[28] = amount
                row[30] = "编号:A;"
                ws.append(row)
            wb.save(sales)
            wb.close()

            process_workbooks(ratio, sales, output)

            result = load_workbook(output, data_only=True)["拆分结果"]
            self.assertEqual(result["AA2"].value, 40)
            self.assertAlmostEqual(result["AD2"].value, 99, places=12)
            self.assertEqual(result["AA3"].value, 30)
            self.assertEqual(result["AA4"].value, 70)
            self.assertAlmostEqual(result["AD3"].value, 29.7, places=12)
            self.assertAlmostEqual(result["AD4"].value, 69.3, places=12)

    def test_conflicting_duplicate_parent_child_is_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            ratio, sales, output = root / "ratio.xlsx", root / "sales.xlsx", root / "output.xlsx"
            make_new_ratio(ratio)
            wb = load_workbook(ratio)
            ws = wb.worksheets[1]
            duplicate = [ws.cell(2, col).value for col in range(1, ws.max_column + 1)]
            duplicate[12] = 41
            ws.append(duplicate)
            wb.save(ratio)
            wb.close()
            make_sales(sales)

            with self.assertRaisesRegex(SplitterError, "执行价格或分摊比例"):
                process_workbooks(ratio, sales, output)

    def test_zero_base_price_uses_source_unit_price(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            ratio, sales, output = root / "ratio.xlsx", root / "sales.xlsx", root / "output.xlsx"
            make_new_ratio(ratio)
            wb = load_workbook(ratio)
            ratio_sheet = wb.worksheets[1]
            ratio_sheet.append(
                ["ZERO", "SKU-ZERO", "零价格子件", "", "", "", "", "", "", "", "件", 1, 0, 0, 1, "否", "否", ""]
            )
            wb.save(ratio)
            wb.close()

            wb = Workbook()
            ws = wb.active
            ws.title = "销售明细"
            ws.append(SALES_HEADERS)
            row = [""] * len(SALES_HEADERS)
            row[1] = "ORDER-ZERO"
            row[10] = "WEB-ZERO"
            row[13] = 25
            row[21] = "SKU-ZERO"
            row[24] = 1
            row[25] = 25
            row[28] = 25
            row[30] = "编号:ZERO;"
            ws.append(row)
            wb.save(sales)
            wb.close()

            process_workbooks(ratio, sales, output)

            result = load_workbook(output, data_only=True)["拆分结果"]
            self.assertEqual(result["AA2"].value, 25)
            self.assertAlmostEqual(result["AD2"].value, 24.75, places=12)

    def test_shuffled_sales_columns_are_matched_by_header_name(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            ratio, sales, output = root / "ratio.xlsx", root / "sales.xlsx", root / "output.xlsx"
            make_ratio(ratio)
            headers = [
                "订单编号",
                "物流单号",
                "网店订单号",
                "应收合计",
                "货品编号",
                "数量",
                "金额",
                "额外字段",
                "单价",
                "备注",
            ]
            wb = Workbook()
            ws = wb.active
            ws.title = "sheetTitle"
            ws.append(headers)
            ws.append(["ORDER-1", "LOG-1", "WEB-1", 100, "SKU1", 1, 50, "保留", 50, ""])
            ws.append(["ORDER-1", "LOG-1", "WEB-1", None, "SKU2", 2, 50, "保留", 25, ""])
            wb.save(sales)
            wb.close()

            stats = process_workbooks(ratio, sales, output)

            self.assertEqual(stats.exceptional_orders, 0)
            result_book = load_workbook(output, data_only=False)
            result = result_book["拆分结果"]
            result_headers = [
                result.cell(1, column).value
                for column in range(1, result.max_column + 1)
            ]
            bundle_col = result_headers.index("组合装单价") + 1
            fee_col = result_headers.index("手续费1%") + 1
            self.assertAlmostEqual(result.cell(2, bundle_col + 3).value, 39.6, places=12)
            self.assertAlmostEqual(result.cell(3, bundle_col + 3).value, 59.4, places=12)
            self.assertAlmostEqual(result.cell(2, fee_col + 2).value, 39.6, places=12)
            self.assertAlmostEqual(result.cell(3, fee_col + 2).value, 59.4, places=12)
            self.assertEqual(result.cell(2, fee_col).value, "=ROUND(G2*1%,2)")
            self.assertEqual(result.cell(2, fee_col - 1).value, "=G2/F2")
            self.assertEqual(result.cell(2, fee_col + 1).value, "=D2-I2")
            self.assertEqual(result["R1"].value, "备注")
            self.assertEqual(result.auto_filter.ref, "A1:R3")
            summary = result_book["透视表"]
            self.assertAlmostEqual(summary["B2"].value, 99, places=12)
            self.assertAlmostEqual(summary["C2"].value, 99, places=12)
            result_book.close()

    def test_three_colleague_column_layouts_are_matched_by_header_name(self):
        layouts = (
            [
                "货品数量",
                "订单编号",
                "网店订单号",
                "物流单号",
                "应收合计",
                "金额",
                "单价",
                "货品编号",
                "数量",
                "备注",
            ],
            [
                "订单编号",
                "物流单号",
                "网店订单号",
                "应收合计",
                "货品编号",
                "数量",
                "单价",
                "备注",
                "金额",
            ],
            [
                "备注",
                "金额",
                "数量",
                "货品编号",
                "应收合计",
                "网店\n订单号",
                "单价",
                "物流单号",
                "订单编号",
            ],
        )
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            ratio = root / "ratio.xlsx"
            make_ratio(ratio)
            for index, headers in enumerate(layouts, start=1):
                with self.subTest(layout=index):
                    sales = root / f"sales-{index}.xlsx"
                    output = root / f"output-{index}.xlsx"
                    wb = Workbook()
                    ws = wb.active
                    ws.title = "sheetTitle"
                    ws.append(headers)
                    values = {
                        "订单编号": "ORDER-1",
                        "物流单号": "LOG-1",
                        "网店订单号": "WEB-1",
                        "应收合计": 100,
                        "货品编号": "SKU1",
                        "数量": 1,
                        "单价": 100,
                        "金额": 100,
                        "备注": "",
                        "货品数量": 1,
                    }
                    ws.append(
                        [
                            values.get(header.replace("\n", "").strip(), "")
                            for header in headers
                        ]
                    )
                    wb.save(sales)
                    wb.close()

                    stats = process_workbooks(ratio, sales, output)

                    self.assertEqual(stats.orders, 1)
                    self.assertEqual(stats.exceptional_orders, 0)
                    result_book = load_workbook(output, data_only=False)
                    summary = result_book["透视表"]
                    self.assertEqual(summary["A2"].value, "WEB-1")
                    self.assertEqual(summary["B2"].value, 99)
                    self.assertEqual(summary["C2"].value, 99)
                    result_book.close()

    def test_zero_receivable_order_is_valid_without_source_amount(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            ratio, sales, output = root / "ratio.xlsx", root / "sales.xlsx", root / "output.xlsx"
            make_ratio(ratio)
            make_sales(sales)
            wb = load_workbook(sales)
            ws = wb["销售明细"]
            for row in range(2, 5):
                ws.cell(row, 14, 0)
                ws.cell(row, 26, 0)
                ws.cell(row, 29, 0)
            wb.save(sales)
            wb.close()

            stats = process_workbooks(ratio, sales, output)

            self.assertEqual(stats.exceptional_orders, 0)
            result = load_workbook(output, data_only=True)["拆分结果"]
            self.assertEqual(sum(result.cell(row, 30).value for row in range(2, 5)), 0)
            self.assertEqual(sum(result.cell(row, 37).value for row in range(2, 5)), 0)

    def test_positive_receivable_all_zero_order_uses_ratio_or_quantity_fallback(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            ratio, sales, output = (
                root / "ratio.xlsx",
                root / "sales.xlsx",
                root / "output.xlsx",
            )
            make_new_ratio(ratio)
            ratio_book = load_workbook(ratio)
            ratio_book["sheet11"]["M2"] = 0
            ratio_book.save(ratio)
            ratio_book.close()

            wb = Workbook()
            ws = wb.active
            ws.title = "销售明细"
            ws.append(SALES_HEADERS)
            for (
                order,
                logistics,
                web_order,
                receivable,
                item,
                quantity,
                note,
            ) in (
                (
                    "ORDER-1",
                    "TRACKING-1",
                    "WEB-1",
                    26,
                    "SKU1",
                    5,
                    "编号:A;",
                ),
                (
                    "ORDER-2",
                    "TRACKING-2",
                    "WEB-2",
                    21,
                    "MISS",
                    1,
                    "编号:UNKNOWN;",
                ),
            ):
                row = [""] * len(SALES_HEADERS)
                row[1] = order
                row[9] = logistics
                row[10] = web_order
                row[13] = receivable
                row[21] = item
                row[24] = quantity
                row[25] = 0
                row[28] = 0
                row[30] = note
                ws.append(row)
            wb.save(sales)
            wb.close()

            stats = process_workbooks(ratio, sales, output)

            self.assertEqual(stats.exceptional_orders, 0)
            result = load_workbook(output, data_only=True)["拆分结果"]
            self.assertEqual(result["AA2"].value, 0)
            self.assertEqual(result["AD2"].value, 26)
            self.assertEqual(result["AK2"].value, 26)
            self.assertEqual(result["AD3"].value, 21)
            self.assertEqual(result["AK3"].value, 21)

    def test_no_fee_mode_uses_receivable_and_writes_only_five_result_columns(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            ratio, sales, output = root / "ratio.xlsx", root / "sales.xlsx", root / "output.xlsx"
            make_ratio(ratio)
            make_sales(sales)

            process_workbooks(ratio, sales, output, charge_fee=False)

            result_book = load_workbook(output, data_only=False)
            result = result_book["拆分结果"]
            headers = [
                result.cell(1, column).value
                for column in range(1, result.max_column + 1)
            ]
            self.assertEqual(result.max_column, len(SALES_HEADERS) + 5)
            self.assertNotIn("单价2", headers)
            self.assertNotIn("手续费1%", headers)
            self.assertNotIn("最终金额", headers)
            bundle_index = headers.index("组合装单价")
            self.assertEqual(
                headers[bundle_index:bundle_index + 4],
                ["组合装单价", "占比", "单价", "金额"],
            )
            self.assertEqual(headers.count("单价"), 2)
            self.assertEqual(headers.count("金额"), 2)
            self.assertEqual(headers.count("摊后金额"), 1)
            split_amount_col = bundle_index + 4
            allocated_col = headers.index("摊后金额") + 1
            self.assertAlmostEqual(
                sum(result.cell(row, split_amount_col).value or 0 for row in range(2, 5)),
                200,
                places=8,
            )
            self.assertAlmostEqual(
                sum(result.cell(row, allocated_col).value or 0 for row in range(2, 5)),
                200,
                places=8,
            )
            summary = result_book["透视表"]
            self.assertEqual(summary["B1"].value, "应收合计")
            self.assertEqual(summary["B2"].value, 200)
            self.assertEqual(summary["C2"].value, 200)
            totals = result_book["汇总"]
            self.assertEqual(totals.max_row, 3)
            self.assertEqual([totals.cell(row, 1).value for row in range(1, 4)], ["应收合计", "摊后金额", "差异"])
            result_book.close()

    def test_no_fee_mode_ignores_external_order_amount(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            ratio, sales, output = root / "ratio.xlsx", root / "sales.xlsx", root / "output.xlsx"
            make_ratio(ratio)
            make_sales(sales)
            add_order_amounts(sales, [("WEB-1", 50)])

            process_workbooks(ratio, sales, output, charge_fee=False)

            result_book = load_workbook(output, data_only=False)
            summary = result_book["透视表"]
            self.assertEqual(summary["B2"].value, 200)
            self.assertEqual(summary["C2"].value, 200)
            result_book.close()

    def test_no_fee_summary_sums_web_order_across_logistics_groups(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            ratio, sales, output = root / "ratio.xlsx", root / "sales.xlsx", root / "output.xlsx"
            make_ratio(ratio)
            wb = Workbook()
            ws = wb.active
            ws.title = "销售明细"
            ws.append(SALES_HEADERS)
            for tracking, receivable in (("TRACKING-A", 60), ("TRACKING-B", 40)):
                row = [""] * len(SALES_HEADERS)
                row[1] = tracking
                row[9] = tracking
                row[10] = "WEB-SHARED"
                row[13] = receivable
                row[21] = "SKU1"
                row[24] = 1
                row[25] = receivable
                row[28] = receivable
                ws.append(row)
            wb.save(sales)
            wb.close()

            process_workbooks(ratio, sales, output, charge_fee=False)

            result_book = load_workbook(output, data_only=False)
            summary = result_book["透视表"]
            self.assertEqual(summary["A2"].value, "WEB-SHARED")
            self.assertEqual(summary["B2"].value, 100)
            self.assertEqual(summary["C2"].value, 100)
            result_book.close()

    def test_no_fee_shared_logistics_allocates_receivable_once_across_web_orders(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            ratio, sales, output = root / "ratio.xlsx", root / "sales.xlsx", root / "output.xlsx"
            make_ratio(ratio)
            wb = Workbook()
            ws = wb.active
            ws.title = "销售明细"
            ws.append(SALES_HEADERS)
            for web_order in ("WEB-A", "WEB-B"):
                row = [""] * len(SALES_HEADERS)
                row[1] = "ORDER-SHARED"
                row[9] = "TRACKING-SHARED"
                row[10] = web_order
                row[13] = 100
                row[21] = "SKU1"
                row[24] = 1
                row[25] = 50
                row[28] = 50
                ws.append(row)
            wb.save(sales)
            wb.close()

            process_workbooks(ratio, sales, output, charge_fee=False)

            result_book = load_workbook(output, data_only=False)
            summary = result_book["透视表"]
            self.assertEqual([summary["B2"].value, summary["B3"].value], [50, 50])
            self.assertEqual([summary["C2"].value, summary["C3"].value], [50, 50])
            result_book.close()

    def test_split_marker_combines_logistics_and_row_ratios_sum_to_one(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            ratio, sales, output = root / "ratio.xlsx", root / "sales.xlsx", root / "output.xlsx"
            make_ratio(ratio)
            wb = Workbook()
            ws = wb.active
            ws.title = "销售明细"
            ws.append(SALES_HEADERS)
            for tracking, code, quantity, amount in (
                ("TRACKING-A", "SKU1", 2, 40),
                ("TRACKING-B", "SKU2", 1, 60),
            ):
                row = [""] * len(SALES_HEADERS)
                row[0] = "驳回,拆分,已成本核算"
                row[1] = f"ORDER-{tracking}"
                row[9] = tracking
                row[10] = "WEB-SPLIT"
                row[13] = 100
                row[21] = code
                row[24] = quantity
                row[25] = amount / quantity
                row[28] = amount
                ws.append(row)
            wb.save(sales)
            wb.close()

            stats = process_workbooks(ratio, sales, output, charge_fee=False)

            self.assertEqual(stats.orders, 1)
            result_book = load_workbook(output, data_only=False)
            result = result_book["拆分结果"]
            self.assertEqual([result["AB2"].value, result["AB3"].value], [0.4, 0.6])
            self.assertEqual(result["AB2"].value + result["AB3"].value, 1)
            self.assertEqual(result["AC2"].value, 20)
            self.assertEqual(result["AH2"].value + result["AH3"].value, 100)
            summary = result_book["透视表"]
            self.assertEqual([summary["B2"].value, summary["C2"].value, summary["D2"].value], [100, 100, 0])
            result_book.close()

    def test_merge_marker_sums_web_receivables_and_marks_pivot_differences(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            ratio, sales, output = root / "ratio.xlsx", root / "sales.xlsx", root / "output.xlsx"
            make_ratio(ratio)
            wb = Workbook()
            ws = wb.active
            ws.title = "销售明细"
            ws.append(SALES_HEADERS)
            for web_order, receivable, code in (
                ("WEB-A", 60, "SKU1"),
                ("WEB-B", 40, "SKU2"),
            ):
                row = [""] * len(SALES_HEADERS)
                row[0] = "有赠品,合并,已成本核算"
                row[1] = "ORDER-MERGED"
                row[9] = "TRACKING-MERGED"
                row[10] = web_order
                row[13] = receivable
                row[21] = code
                row[24] = 1
                row[25] = 50
                row[28] = 50
                ws.append(row)
            wb.save(sales)
            wb.close()

            stats = process_workbooks(ratio, sales, output, charge_fee=False)

            self.assertEqual(stats.orders, 1)
            self.assertEqual(stats.exceptional_orders, 0)
            result_book = load_workbook(output, data_only=False)
            result = result_book["拆分结果"]
            self.assertEqual(result["AB2"].value + result["AB3"].value, 1)
            self.assertEqual(result["AH2"].value + result["AH3"].value, 100)
            summary = result_book["透视表"]
            self.assertEqual([summary["B2"].value, summary["B3"].value], [60, 40])
            self.assertEqual([summary["C2"].value, summary["C3"].value], [40, 60])
            self.assertEqual([summary["D2"].value, summary["D3"].value], [-20, 20])
            self.assertEqual(summary["D2"].fill.fgColor.rgb, "00FCE8E6")
            self.assertEqual(summary["D3"].fill.fgColor.rgb, "00FCE8E6")
            self.assertEqual([summary["B4"].value, summary["C4"].value, summary["D4"].value], [100, 100, 0])
            result_book.close()

    def test_split_and_merge_markers_form_one_transitive_group(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            ratio, sales, output = root / "ratio.xlsx", root / "sales.xlsx", root / "output.xlsx"
            make_ratio(ratio)
            wb = Workbook()
            ws = wb.active
            ws.title = "销售明细"
            ws.append(SALES_HEADERS)
            for marker, tracking, web_order, receivable, code in (
                ("拆分,已成本核算", "TRACKING-A", "WEB-A", 60, "SKU1"),
                ("拆分,合并,已成本核算", "TRACKING-B", "WEB-A", 60, "SKU2"),
                ("合并,已成本核算", "TRACKING-B", "WEB-B", 40, "SKU1"),
            ):
                row = [""] * len(SALES_HEADERS)
                row[0] = marker
                row[1] = f"ORDER-{tracking}-{web_order}"
                row[9] = tracking
                row[10] = web_order
                row[13] = receivable
                row[21] = code
                row[24] = 1
                row[25] = 50
                row[28] = 50
                ws.append(row)
            wb.save(sales)
            wb.close()

            stats = process_workbooks(ratio, sales, output, charge_fee=False)

            self.assertEqual(stats.orders, 1)
            self.assertEqual(stats.exceptional_orders, 0)
            result_book = load_workbook(output, data_only=False)
            result = result_book["拆分结果"]
            self.assertEqual(
                sum(result.cell(row, 28).value for row in range(2, 5)),
                1,
            )
            self.assertEqual(
                sum(result.cell(row, 34).value for row in range(2, 5)),
                100,
            )
            result_book.close()

    def test_source_files_are_not_overwritten(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            ratio, sales = root / "ratio.xlsx", root / "sales.xlsx"
            make_ratio(ratio)
            make_sales(sales)
            with self.assertRaises(SplitterError):
                process_workbooks(ratio, sales, sales)


if __name__ == "__main__":
    unittest.main()
