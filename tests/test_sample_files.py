from __future__ import annotations

from collections import OrderedDict
import os
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook

from splitter import (
    SALES_HEADERS,
    _find_sheet,
    _header_map,
    _sales_columns,
    process_workbooks,
)


RATIO = Path(os.environ.get("BUNDLE_SPLITTER_RATIO_SAMPLE", "__missing_ratio_sample__.xlsx"))
SALES = Path(os.environ.get("BUNDLE_SPLITTER_SALES_SAMPLE", "__missing_sales_sample__.xlsx"))


def _group_key(values: tuple, columns, row: int) -> tuple[str, str, int]:
    logistics = str(values[columns.logistics_number - 1] or "").strip()
    order = str(values[columns.order_number - 1] or "").strip()
    web_order = str(values[columns.web_order_number - 1] or "").strip()
    if logistics:
        return ("物流单号", logistics, 0)
    if order:
        return ("订单编号", order, 0)
    if web_order:
        return ("网店订单号", web_order, 0)
    return ("单行", "", row)


@unittest.skipUnless(RATIO.exists() and SALES.exists(), "样例 Excel 不在当前电脑")
class RealSampleTests(unittest.TestCase):
    def _source_receivables(self):
        original = load_workbook(SALES, data_only=True, read_only=True)
        try:
            source, header_row = _find_sheet(original, SALES_HEADERS)
            self.assertIsNotNone(source)
            source_headers = _header_map(source, header_row)
            receivable_col = source_headers["应收合计"]
            receivables = [
                values[0]
                for values in source.iter_rows(
                    min_row=header_row + 1,
                    min_col=receivable_col,
                    max_col=receivable_col,
                    values_only=True,
                )
            ]
            return source.title, header_row, receivable_col, receivables
        finally:
            original.close()

    def test_fee_mode_preserves_source_and_balances_result_rows(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "sample-output.xlsx"
            source_title, header_row, receivable_col, original_receivables = (
                self._source_receivables()
            )

            stats = process_workbooks(RATIO, SALES, output)

            self.assertGreater(stats.matched_rows, 0)
            self.assertEqual(stats.exceptional_orders, 0)
            wb = load_workbook(output, data_only=True, read_only=True)
            try:
                output_receivables = [
                    values[0]
                    for values in wb[source_title].iter_rows(
                        min_row=header_row + 1,
                        min_col=receivable_col,
                        max_col=receivable_col,
                        values_only=True,
                    )
                ]
                self.assertEqual(output_receivables, original_receivables)

                result = wb["拆分结果"]
                result_headers = _header_map(result, header_row)
                bundle_col = result_headers["组合装单价"]
                split_amount_col = bundle_col + 3
                allocated_col = result_headers["摊后金额"]
                for row in result.iter_rows(min_row=header_row + 1, values_only=True):
                    self.assertAlmostEqual(
                        float(row[split_amount_col - 1] or 0),
                        float(row[allocated_col - 1] or 0),
                        places=8,
                    )

            finally:
                wb.close()

    def test_no_fee_mode_balances_every_group_to_receivable(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "sample-output-no-fee.xlsx"

            stats = process_workbooks(RATIO, SALES, output, charge_fee=False)

            self.assertGreater(stats.matched_rows, 0)
            self.assertEqual(stats.exceptional_orders, 0)
            wb = load_workbook(output, data_only=True, read_only=True)
            try:
                result = wb["拆分结果"]
                columns = _sales_columns(result, 1)
                headers = _header_map(result, 1)
                split_amount_col = headers["组合装单价"] + 3
                allocated_col = headers["摊后金额"]
                self.assertNotIn("手续费1%", headers)
                self.assertNotIn("最终金额", headers)
                self.assertNotIn("单价2", headers)

                groups: OrderedDict[tuple[str, str, int], list[tuple]] = OrderedDict()
                for row_number, values in enumerate(
                    result.iter_rows(min_row=2, values_only=True),
                    start=2,
                ):
                    if all(value is None for value in values):
                        continue
                    groups.setdefault(
                        _group_key(values, columns, row_number),
                        [],
                    ).append(values)

                for rows in groups.values():
                    receivables = [
                        float(values[columns.receivable - 1])
                        for values in rows
                        if values[columns.receivable - 1] not in (None, "")
                    ]
                    self.assertTrue(receivables)
                    target = round(receivables[0], 2)
                    self.assertAlmostEqual(
                        sum(float(values[split_amount_col - 1] or 0) for values in rows),
                        target,
                        places=8,
                    )
                    self.assertAlmostEqual(
                        sum(float(values[allocated_col - 1] or 0) for values in rows),
                        target,
                        places=8,
                    )

                summary = wb["透视表"]
                receivable_total = sum(
                    float(summary.cell(row, 2).value or 0)
                    for row in range(2, summary.max_row)
                )
                allocated_total = sum(
                    float(summary.cell(row, 3).value or 0)
                    for row in range(2, summary.max_row)
                )
                self.assertAlmostEqual(receivable_total, allocated_total, places=8)
            finally:
                wb.close()


if __name__ == "__main__":
    unittest.main()
